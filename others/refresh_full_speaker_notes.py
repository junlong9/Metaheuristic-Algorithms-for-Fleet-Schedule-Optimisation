"""Refresh speaker notes in the Milestones deck to full prose scripts."""

from pathlib import Path

from pptx import Presentation

DECK = Path(__file__).resolve().parents[1] / "presentations" / "Iteration 1 Milestones.pptx"
OUT = DECK.with_name("Iteration 1 Milestones - scripted.pptx")

SCRIPTS = {
    1: """Thanks everyone, and welcome if you are joining us for the first time.

I will start with a bit of context so we are all on the same page. This project is about scheduling heavy maintenance for a fleet of aircraft. Each aircraft needs a sequence of planned checks over a long forecast horizon — A, B, C, and D checks — and those checks have legal windows for when they can start. The operational problem is that if too many aircraft go into maintenance at the same time, fleet availability drops, and that hurts the ability to fly the mission.

So the goal is not to cancel maintenance. The goal is to stagger it. We shift when each check starts, still inside the allowed windows, so the dips in availability are as shallow as possible and the day-to-day curve is as smooth as we can make it.

The way we score a schedule reflects that. First we try to raise the minimum number of aircraft available on the worst day in the forecast. If two schedules have the same floor, we prefer the one with less variation. If those are still tied, we look at average availability. And alongside quality, the method has to finish in a practical amount of time, because once the dataset gets large, a slow optimiser is hard to use.

What I am going to walk through today is Iteration 1: how we started with a simple proof of concept, tried a few search strategies on a heavier dataset, and ended up with a priority-based method that improves the schedule and runs much faster than the earlier approaches.""",
    2: """This is Iteration 1, results and review. The short version of the outcome is that we moved from exhaustive neighbourhood search to priority-based selection, and that is what made the optimiser practical.

On the heavy dataset, the priority method lifted minimum availability from fifteen aircraft to eighteen, and it reduced the standard deviation from about 1.96 down to 1.44. Average availability barely changed. That is actually fine for us, because the objective is deliberately lexicographic. We care about the worst day first. A schedule that protects the floor is more valuable than one that only lifts the average.

I will show the path we took to get there, milestone by milestone.""",
    3: """There are four milestones in this iteration.

Milestone one was a simple fleet dataset. The point there was to prove that the optimisation objective and the feasibility rules worked at all.

Then we moved to a heavy multi-year forecast with twenty-four aircraft and two hundred and seventy-eight maintenance events. On that dataset we tried three search strategies. First, full neighbourhood search. Second, a shortcut that randomly sampled two hundred neighbours. Third, priority-based selection, which is where we landed.

The story is fairly straightforward. First make it work. Then see what breaks when you scale. Then stop searching blindly and start searching where the availability problem actually is.""",
    4: """Milestone one used a small random fleet plan with eighty maintenance events. The optimiser shifts event start times inside feasible windows, rejects overlapping maintenance on the same aircraft, and scores each schedule with that same order I mentioned: maximise minimum availability, then minimise variation, then maximise the average.

We used variable neighbourhood search, so we try larger shifts first and then move to finer ones. That helps escape poor local solutions without jumping straight into tiny adjustments.

The important result from this milestone was not that we had an industrial-ready system. It was that we could stagger a schedule under hard constraints and improve that lexicographic score. One thing I want to flag early: please do not compare runtimes from this simple dataset with the heavy-dataset timings later. They are different problem sizes, and treating them as a direct speed comparison would be misleading.""",
    5: """This is the dataset that really stressed the algorithms.

Maintenance follows the repeating cycle A, B, A, C, A, B, A, D. The forecast is at least twelve years long, which is about one and a half maintenance cycles. There are twenty-four aircraft, each with roughly eleven to twelve events in the forecast, for two hundred and seventy-eight events in total.

Once you are at that scale, every candidate move becomes more expensive. Deep-copying schedules, checking feasibility, and recomputing availability add up quickly if you do it for every aircraft, every event, and every shift. The next three milestones all use this same heavy dataset, so those comparisons are fair. The simple dataset stays in the story as the proof of concept, not as the benchmark for speed.""",
    6: """Full neighbourhood search is the thorough baseline. For every aircraft and every event, we try the shift sizes: plus or minus twenty-eight days, then fourteen, seven, and three. Each candidate is created by deep-copying the whole fleet schedule. Then we check feasibility and score the full schedule.

That approach does improve quality, but the cost grows with fleet size times number of events times number of shifts. On this heavy dataset it took about fifteen minutes — around nine hundred and twenty seconds. That is acceptable if you are running a one-off experiment. It is not ideal if you want to iterate, debug, or eventually plug the optimiser into a planning workflow.

So the next question was obvious: can we keep most of the quality while evaluating fewer candidates?""",
    7: """Milestone three uses the same heavy dataset and the same shift sets, but with a sampling cutoff. If the neighbourhood contains more than two hundred feasible neighbours, we randomly keep two hundred and only score those.

That did reduce runtime. It came down to about seven and a half minutes, around four hundred and fifty-one seconds. So sampling helped. But it also introduced a real trade-off. Because the selection is random, you can miss the best improving move in that neighbourhood. Unless you fix a seed, the run is also non-deterministic. And importantly, we were still deep-copying full schedules and fully revalidating candidates, so a large part of the computational cost was still sitting there.

This was a useful intermediate step. It showed that thinning the search can help, but random thinning is not really a search strategy. It is a compromise. That motivated a smarter, deterministic way to choose which candidates to evaluate.""",
    8: """This is the method I want to spend the most time on, because it is the main technical result of Iteration 1.

Priority-based selection changes the question we ask. Instead of asking what happens if we try every feasible move, or what happens if we try a random subset, we ask which maintenance events are actually hurting the worst days in the availability curve.

Every event gets a priority equal to the lowest fleet availability on any day that event is in the hangar. So if an event overlaps a day when only fifteen aircraft are available, that event is priority fifteen. Those events are the ones we search first, because they are the ones sitting on the availability floor.

When we move one check, we do not leave the rest of that aircraft's plan sitting in a potentially broken state. We rebuild the later events using the A, B, C, D maintenance rules. Larger checks reset the smaller ones in the cycle, due windows are recalculated, and starts are kept where they are still legal or clamped into the new legal window. If a rebuild cannot produce a valid downstream plan, we reject that candidate.

We also changed how scoring works. We do not deep-copy the entire fleet for every trial. We rebuild one aircraft, then update availability incrementally: remove that aircraft's old hangar intervals from the stored availability array, add the rebuilt intervals, and score once. That is much cheaper than recomputing availability from scratch for every neighbour.

On the heavy dataset, that gave us a better schedule — minimum availability from fifteen to eighteen, standard deviation from about 1.96 to 1.44 — and a median runtime of about fifteen seconds. So the quality improved, and the method became practical.""",
    9: """I want to walk through the mechanics of those priority levels, because this is usually where the detailed questions come up.

Take the starting heavy schedule. The global minimum availability is fifteen. First, we score every event. For each event we look at fleet availability on every day that event is in the hangar, and the priority is the minimum of those days. For example, if an event covers days with availability eighteen, fifteen, sixteen, and so on, its priority is fifteen.

Second, we split the events into exclusive buckets by that priority. In our starting schedule we had twenty events at level fifteen, ten at sixteen, nine at seventeen, and larger groups at higher levels. An event sits in only one bucket: the bucket for its own worst day. So the twenty events at fifteen and the ten at sixteen are different events. They are not stages of the same repair ladder.

Third, we search worst first. For a shift set like plus and minus twenty-eight, level fifteen is about twenty events times two shifts, so around forty candidates. If any of those moves improve the score, we stop at that level for the iteration and take the best improving move. We only expand to level sixteen, then seventeen, and so on if the worse levels find nothing.

That is why the method is both targeted and fast. We are not randomly discarding most of the neighbourhood. We are concentrating effort on the events that touch the availability floor, and we only widen the search when we have to.

One detail that sometimes confuses people: on a day when availability is fifteen, only nine aircraft are actually in maintenance, because twenty-four minus fifteen is nine. That does not mean there are only nine priority-fifteen events. Priority is assigned across the whole timeline. Many different events can each brush a fifteen-day dip at some point, so twenty events can carry priority fifteen even though only nine aircraft are down on any single trough day.""",
    10: """Here are the measured results on the same heavy dataset.

For schedule quality under the priority method, minimum availability went from fifteen to eighteen, standard deviation from 1.96 to 1.44, and average availability stayed almost flat at about 21.76 to 21.75.

For runtime, full neighbourhood search took about nine hundred and twenty seconds, which is roughly fifteen minutes. The two-hundred random sample method took about four hundred and fifty-one seconds, which is roughly seven and a half minutes. Priority selection took about fifteen point four seconds as the median of three runs.

So we improved the worst-case availability and the variation, and we cut runtime by about sixty times compared with full neighbourhood search, and by about thirty times compared with random sampling. That is the Iteration 1 result we wanted to demonstrate: better schedules, and a method you can actually run repeatedly.""",
    11: """Looking ahead to Semester Two and Iteration 2, the plan is to investigate a simulated annealing and variable neighbourhood search hybrid, bring in additional operational and logistics constraints, analyse hyperparameters more carefully, and test whether the initial stagger of the schedule changes the final outcome. In the nearer term that also includes thinking about windows where we maximise availability and allow surrounding events to move, while controlling standard deviation inside those windows.

That is Iteration 1. I am happy to take questions, especially on the priority-based method.""",
}


def set_notes(slide, text):
    tf = slide.notes_slide.notes_text_frame
    if tf is None:
        raise RuntimeError("Slide has no notes text frame")
    for paragraph in list(tf.paragraphs)[1:]:
        paragraph._p.getparent().remove(paragraph._p)
    p0 = tf.paragraphs[0]
    for run in list(p0.runs):
        run._r.getparent().remove(run._r)
    run = p0.add_run()
    run.text = text


def main():
    path = DECK if DECK.exists() else ALT
    prs = Presentation(str(path))
    print(f"Updating notes in {path.name} ({len(prs.slides)} slides)")
    for index, text in SCRIPTS.items():
        set_notes(prs.slides[index - 1], text)
        print(f"  slide {index}")
    try:
        prs.save(str(path))
        print(f"Saved {path}")
    except PermissionError:
        out = path.with_name(path.stem + " - full script.pptx")
        prs.save(str(out))
        print(f"Locked; saved {out}")


if __name__ == "__main__":
    main()
