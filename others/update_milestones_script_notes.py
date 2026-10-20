"""Insert project-intro slide and embed speaking notes into Iteration 1 Milestones.pptx."""

from copy import deepcopy
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_AUTO_SHAPE_TYPE
from pptx.enum.text import PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt

SRC = Path(__file__).resolve().parents[1] / "presentations" / "Iteration 1 Milestones.pptx"
ALT = Path(__file__).resolve().parents[1] / "presentations" / "Iteration 1 Milestones with intro.pptx"

BLACK = RGBColor(0, 0, 0)
WHITE = RGBColor(255, 255, 255)
PALE = RGBColor(0xF0, 0xF0, 0xF0)
LINE = RGBColor(0xB0, 0xB0, 0xB0)


def set_run(run, size=16, bold=False, color=BLACK):
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    run.font.name = "Calibri"


def add_box(slide, left, top, width, height, text, *, size=16, bold=False, color=BLACK):
    box = slide.shapes.add_textbox(left, top, width, height)
    tf = box.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.LEFT
    run = p.add_run()
    run.text = text
    set_run(run, size=size, bold=bold, color=color)
    return box


def add_metric(slide, top, label, value):
    card = slide.shapes.add_shape(
        MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE,
        Inches(8.35),
        top,
        Inches(4.4),
        Inches(1.5),
    )
    card.fill.solid()
    card.fill.fore_color.rgb = PALE
    card.line.color.rgb = BLACK
    add_box(
        slide,
        Inches(8.55),
        top + Inches(0.2),
        Inches(4.0),
        Inches(0.35),
        label,
        size=12,
        bold=True,
    )
    add_box(
        slide,
        Inches(8.55),
        top + Inches(0.6),
        Inches(4.0),
        Inches(0.7),
        value,
        size=22,
        bold=True,
    )


SCRIPTS = {
    1: """SCRIPT (~1:30)

Thanks everyone — and welcome if you are joining for the first time.

We have a fleet of aircraft that each need planned heavy maintenance over a long horizon — A, B, C, and D checks on a repeating cycle. If too many aircraft go into maintenance at the same time, fleet availability drops.

So the job is not "do less maintenance." It is to stagger the work: shift when each check starts, inside legal windows, so the dips in availability are as shallow and as smooth as we can make them.

What we measure: on the worst day, how many aircraft are still available? How bumpy is the curve? And can the optimiser answer that without taking forever?

This review is how we went from a simple proof-of-concept, through methods that were too slow, to priority-based search that gets better schedules in seconds.

BRIDGE: I will start with where we ended up, then walk through each milestone.""",
    2: """SCRIPT (~0:45)

This is Iteration 1 — results and review. Short version: we moved from exhaustive neighbourhood search to priority-based selection, and that made the optimiser practical.

On the heavy dataset, priority lifted minimum availability from 15 to 18 aircraft, and cut standard deviation from about 1.96 to 1.44. Average barely moved — which is fine, because we care about the worst day first.

BRIDGE: Here is the path we took.""",
    3: """SCRIPT (~1:00)

Four milestones.

Milestone 1: simple fleet — prove the objective and feasibility rules.
Then the heavy forecast: 24 aircraft, 278 events. On that data we tried full neighbourhood search, then 200 random samples, then priority-based selection.

Pattern: make it work, make it scale, then make the search smarter instead of just thinner.

BRIDGE: Starting with the simple case.""",
    4: """SCRIPT (~1:00)

Milestone 1 used a small random fleet plan — eighty events. We shift start times inside allowed windows, reject overlaps, and score with a clear order: raise minimum availability first, then reduce variation, then average.

Variable neighbourhood search — bigger shifts first, then finer ones. The point was proof we could stagger the schedule under rules. Do not compare runtimes here with the heavy-dataset numbers later.

BRIDGE: Then we moved to a realistic heavy forecast.""",
    5: """SCRIPT (~1:10)

This is the dataset that stressed the algorithms.

Cycle A, B, A, C, A, B, A, D. Forecast at least twelve years — about one and a half cycles — twenty-four aircraft, roughly eleven to twelve events each, two hundred seventy-eight events total.

At that scale naive search stops being practical. The next three milestones all use this same heavy dataset, so those comparisons are fair.

BRIDGE: First attempt — evaluate the full neighbourhood.""",
    6: """SCRIPT (~1:10)

Full neighbourhood is the thorough baseline. For every aircraft and event, try the shift sizes — plus or minus 28, 14, 7, 3. Each candidate is a deep copy of the whole schedule, then full feasibility and scoring.

Quality improved, but cost scaled with fleet size times events times shifts. On this set it took about fifteen minutes — around 920 seconds. Fine for an experiment; not fine for iteration.

BRIDGE: So we tried cutting the work with random sampling.""",
    7: """SCRIPT (~1:00)

Same dataset and shift sets — but if we generate more than 200 neighbours, we randomly keep 200 and only score those.

That roughly halved runtime to about seven and a half minutes. Catch: you might miss the best improving move, and selection is non-deterministic unless you fix a seed. We still deep-copied candidates, so a lot of cost remained.

Useful intermediate step — sampling helps, but random is not a real search strategy.

BRIDGE: That leads to priority-based selection.""",
    8: """SCRIPT (~2:00) — go slower

Priority flips the question. Instead of try everything or try a random subset, we ask: which maintenance events are hurting the worst days?

Each event gets a priority equal to the lowest fleet availability on any day that event is in the hangar. Events on a day-15 trough are priority 15 — searched first.

When we move one check, we rebuild later events on that aircraft with A/B/C/D rules so due windows stay legal.

We do not deep-copy the whole fleet for every trial. Rebuild one aircraft, patch availability incrementally, score once.

Result: min 15 to 18, std about 1.96 to 1.44, median runtime about 15 seconds.

BRIDGE: Next slide is how the priority buckets work.""",
    9: """SCRIPT (~2:00)

Starting heavy schedule — global min availability is 15.

One: for every event, look at availability across its hangar days. Priority is the minimum of those days.

Two: exclusive buckets. In our start we had 20 events at level 15, 10 at 16, 9 at 17, and so on. An event sits in one bucket only.

Three: search worst first. Level 15 with plus/minus 28 is about 20 events times 2 shifts — around 40 candidates. If any improve, take the best and stop. Only open level 16 if 15 finds nothing.

Not randomly thinning the search — concentrating on events that touch the availability floor.

BRIDGE: Numbers side by side.""",
    10: """SCRIPT (~1:15)

Same heavy dataset for all three.

Quality for priority: min 15 to 18, std 1.96 to 1.44, average basically flat.

Runtime: full neighbourhood about 920 seconds, random 200 about 451, priority about 15.4 seconds median.

Kept — actually improved — quality, and about sixty times faster than full search. That is the Iteration 1 result.

BRIDGE: What is next.""",
    11: """SCRIPT (~0:50)

Semester two — Iteration 2 — simulated annealing plus VNS hybrid, tighter operational and logistics constraints, hyperparameter analysis, and whether the initial stagger changes the outcome.

Happy to take questions — especially on the priority method.

Q&A CHEAT SHEET (priority):
- Faster than random 200 because no full-fleet deep copies; rebuild one aircraft + incremental availability; search worst bucket first.
- Priority 15 is an availability level, not a sample size of 15 or 20.
- Buckets are exclusive, not nested.
- Rebuild costs more than a naive single-date shift, but still far cheaper than deep-copying the fleet, and keeps schedules rule-legal.
- Deterministic given the same inputs.
See presentation_speaker_script.md for fuller Q&A answers.""",
}


def main():
    prs = Presentation(str(SRC))
    print(f"Opened {len(prs.slides)} slides")

    blank = prs.slide_layouts[6]
    slide = prs.slides.add_slide(blank)

    bar = slide.shapes.add_shape(
        MSO_AUTO_SHAPE_TYPE.RECTANGLE,
        0,
        0,
        prs.slide_width,
        Inches(0.12),
    )
    bar.fill.solid()
    bar.fill.fore_color.rgb = BLACK
    bar.line.fill.background()

    add_box(
        slide,
        Inches(0.55),
        Inches(0.35),
        Inches(12.0),
        Inches(0.3),
        "PROJECT CONTEXT",
        size=12,
        bold=True,
    )
    add_box(
        slide,
        Inches(0.55),
        Inches(0.7),
        Inches(12.2),
        Inches(0.7),
        "What this project is about",
        size=28,
        bold=True,
    )

    card = slide.shapes.add_shape(
        MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE,
        Inches(0.55),
        Inches(1.55),
        Inches(7.5),
        Inches(5.2),
    )
    card.fill.solid()
    card.fill.fore_color.rgb = WHITE
    card.line.color.rgb = LINE

    bullets = [
        "A fleet of aircraft each need planned heavy maintenance (A/B/C/D checks) over a long horizon.",
        "If too many go into maintenance together, fleet availability drops.",
        "The goal is to stagger start times inside legal windows to protect the worst days.",
        "We score schedules by: raise minimum availability first, then reduce variation, then average.",
        "Iteration 1 compares four optimiser designs on that problem — ending in priority-based search.",
    ]
    box = slide.shapes.add_textbox(Inches(0.85), Inches(1.8), Inches(7.0), Inches(4.7))
    tf = box.text_frame
    tf.word_wrap = True
    for i, item in enumerate(bullets):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.space_after = Pt(10)
        run = p.add_run()
        run.text = item
        set_run(run, size=15)

    add_metric(slide, Inches(1.55), "OPERATIONAL IDEA", "Stagger maintenance")
    add_metric(slide, Inches(3.3), "PRIMARY METRIC", "Min availability")
    add_metric(slide, Inches(5.05), "THIS REVIEW", "4 algorithm milestones")

    # Move newly added slide to front
    sld_id_lst = prs.slides._sldIdLst
    children = list(sld_id_lst)
    last = children[-1]
    sld_id_lst.remove(last)
    sld_id_lst.insert(0, last)
    print(f"Intro moved to front; total {len(prs.slides)} slides")

    def ensure_notes_frame(slide, template_slide):
        """Blank layouts in this deck create empty notes slides; clone placeholders if needed."""
        notes_slide = slide.notes_slide
        if notes_slide.notes_text_frame is not None:
            return notes_slide.notes_text_frame

        template_notes = template_slide.notes_slide
        # Replace empty notes cSld shapes with a clone of the template notes shapes
        notes_cSld = notes_slide._element.cSld
        spTree = notes_cSld.find(qn("p:spTree"))
        template_spTree = template_notes._element.cSld.find(qn("p:spTree"))

        # Keep nvGrpSpPr / grpSpPr (first two children), replace shape children
        keep = list(spTree)[:2]
        for child in list(spTree):
            spTree.remove(child)
        for child in keep:
            spTree.append(child)
        for child in list(template_spTree)[2:]:
            spTree.append(deepcopy(child))

        tf = notes_slide.notes_text_frame
        if tf is None:
            raise RuntimeError("Failed to clone notes placeholders")
        return tf

    def set_notes(slide, text, template_slide):
        tf = ensure_notes_frame(slide, template_slide)
        for paragraph in list(tf.paragraphs)[1:]:
            paragraph._p.getparent().remove(paragraph._p)
        p0 = tf.paragraphs[0]
        for run in list(p0.runs):
            run._r.getparent().remove(run._r)
        run = p0.add_run()
        run.text = text

    template = prs.slides[1]  # original title slide after intro insert
    for index, text in SCRIPTS.items():
        set_notes(prs.slides[index - 1], text, template)
        print(f"Notes set on slide {index}")

    try:
        prs.save(str(SRC))
        print(f"Saved: {SRC}")
    except PermissionError:
        prs.save(str(ALT))
        print(f"Original locked in PowerPoint. Saved: {ALT}")
        print("Close the open deck, then replace or copy notes from the new file.")


if __name__ == "__main__":
    main()
