from pathlib import Path
import csv
import json
import sys
import tempfile
import unittest

import numpy as np
import pandas as pd
from docx import Document

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))

from analysis import vn_hyperparameter_analysis as study


class MatrixTests(unittest.TestCase):
    def test_strategy_definitions_match_the_approved_matrix(self):
        actual = {
            strategy.key: strategy.shift_sets
            for strategy in study.STRATEGIES
        }
        self.assertEqual(
            actual,
            {
                "historical_half_step": (
                    (-28, -14, 14, 28),
                    (-14, -7, 7, 14),
                    (-7, -3, 3, 7),
                    (-3, -1, 1, 3),
                ),
                "live_wide_to_fine": (
                    (-56, 56),
                    (-42, 42),
                    (-28, 28),
                    (-14, 14),
                    (-7, 7),
                    (-3, 3),
                    (-1, 1),
                ),
                "dense_pairs": (
                    (-28, 28),
                    (-21, 21),
                    (-14, 14),
                    (-7, 7),
                    (-3, 3),
                    (-1, 1),
                ),
                "half_step_plus_42": (
                    (-42, 42),
                    (-28, -14, 14, 28),
                    (-14, -7, 7, 14),
                    (-7, -3, 3, 7),
                    (-3, -1, 1, 3),
                ),
            },
        )
        self.assertEqual(study.BUDGETS, (200, 400, 800))

    def test_trial_matrix_has_each_configuration_once(self):
        matrix = study.build_trial_matrix()
        keys = [
            (dataset, strategy.key, budget)
            for dataset, strategy, budget in matrix
        ]
        self.assertEqual(len(keys), 36)
        self.assertEqual(len(set(keys)), 36)
        self.assertEqual(
            {dataset for dataset, _, _ in keys},
            {"original", "moderate", "wide"},
        )


class InputLoadingTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.inputs = study.load_study_inputs(ROOT)

    def test_inputs_use_the_common_reporting_window(self):
        inputs = self.inputs
        self.assertEqual(
            inputs.origin + pd.Timedelta(days=inputs.scoring_start),
            pd.Timestamp("2026-01-21"),
        )
        self.assertEqual(
            inputs.origin + pd.Timedelta(days=inputs.scoring_end),
            pd.Timestamp("2038-01-18"),
        )
        self.assertEqual(
            tuple(dataset.key for dataset in inputs.datasets),
            ("original", "moderate", "wide"),
        )

    def test_stagger_inputs_keep_saved_opening_state(self):
        for dataset in self.inputs.datasets[1:]:
            for events in dataset.schedule.values():
                self.assertIn("_origin", events[0])
                self.assertIn("_opening", events[0])

    def test_every_initial_schedule_is_rule_compliant(self):
        for dataset in self.inputs.datasets:
            self.assertTrue(
                study.optimiser.is_rule_compliant_schedule(
                    dataset.schedule,
                    self.inputs.horizon,
                )
            )
        self.assertEqual(
            set(self.inputs.input_hashes),
            {"original", "moderate", "wide", "opening_state", "production"},
        )

    def test_missing_inputs_name_the_missing_source(self):
        with tempfile.TemporaryDirectory() as directory:
            with self.assertRaisesRegex(
                FileNotFoundError,
                "opening_state: missing input",
            ):
                study.load_study_inputs(Path(directory))


class TrialTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.inputs = study.load_study_inputs(ROOT)

    def test_independent_availability_matches_production_calculation(self):
        dataset = self.inputs.datasets[0]
        expected = study.optimiser.compute_availability(
            dataset.schedule,
            self.inputs.horizon,
        )
        actual = study.independent_availability(
            dataset.schedule,
            self.inputs.horizon,
        )
        np.testing.assert_array_equal(actual, expected)

    def test_schedule_hash_ignores_dictionary_insertion_order(self):
        first = {
            "B": [{
                "check": "A",
                "start": 1,
                "duration": 14,
                "earliest_start": 0,
                "latest_start": 2,
            }],
            "A": [{
                "check": "A",
                "start": 2,
                "duration": 14,
                "earliest_start": 0,
                "latest_start": 3,
            }],
        }
        second = {"A": first["A"], "B": first["B"]}
        self.assertEqual(
            study.schedule_sha256(first),
            study.schedule_sha256(second),
        )

    def test_one_iteration_trial_is_real_deterministic_and_compliant(self):
        dataset = self.inputs.datasets[0]
        result = study.run_trial(
            dataset,
            study.STRATEGIES[0],
            1,
            horizon=self.inputs.horizon,
            scoring_start=self.inputs.scoring_start,
            scoring_end=self.inputs.scoring_end,
            clock=iter((10.0, 12.5)).__next__,
        )
        self.assertEqual(result["trial_key"], "original:historical_half_step:1")
        self.assertEqual(result["elapsed_seconds"], 2.5)
        self.assertLessEqual(result["accepted_moves"], 1)
        self.assertTrue(result["rule_compliant"])
        self.assertEqual(
            result["stop_reason"],
            "iteration_cap" if result["accepted_moves"] else "local_exhaustion",
        )
        scores = [tuple(item["score"]) for item in result["trace"]]
        self.assertTrue(
            all(left < right for left, right in zip(scores, scores[1:]))
        )


def fake_trial(dataset, strategy, max_iters, **_kwargs):
    strategy_index = next(
        index
        for index, item in enumerate(study.STRATEGIES)
        if item.key == strategy.key
    )
    score = [
        20.0 - strategy_index,
        -(1.0 + max_iters / 10000),
        21.0,
    ]
    return {
        "trial_key": f"{dataset.key}:{strategy.key}:{max_iters}",
        "dataset": dataset.key,
        "dataset_label": dataset.label,
        "strategy_key": strategy.key,
        "strategy_label": strategy.label,
        "shift_sets": [list(level) for level in strategy.shift_sets],
        "max_iters": max_iters,
        "min_availability": int(score[0]),
        "availability_std": -score[1],
        "mean_availability": score[2],
        "lexicographic_score": score,
        "accepted_moves": 2,
        "outer_passes": 3,
        "stop_reason": "local_exhaustion",
        "elapsed_seconds": float(max_iters),
        "rule_compliant": True,
        "endpoint_sha256": strategy.key,
        "trace": [{"iteration": 0, "score": score}],
    }


class StudyOutputTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.inputs = study.load_study_inputs(ROOT)

    def test_run_study_calls_each_of_36_trials_once(self):
        calls = []

        def recording_runner(*args, **kwargs):
            calls.append((args[0].key, args[1].key, args[2]))
            return fake_trial(*args, **kwargs)

        payload = study.run_study(self.inputs, recording_runner)
        self.assertEqual(len(calls), 36)
        self.assertEqual(len(set(calls)), 36)
        self.assertEqual(len(payload["trials"]), 36)

    def test_exact_score_ties_share_competition_rank(self):
        trials = [
            {"dataset": "original", "lexicographic_score": [18, -1.2, 21]},
            {"dataset": "original", "lexicographic_score": [18, -1.2, 21]},
            {"dataset": "original", "lexicographic_score": [17, -1.0, 22]},
        ]
        study.assign_dataset_ranks(trials)
        self.assertEqual([item["dataset_rank"] for item in trials], [1, 1, 3])

    def test_json_and_csv_cover_the_same_36_trials(self):
        payload = study.run_study(self.inputs, fake_trial)
        with tempfile.TemporaryDirectory() as directory:
            paths = study.write_machine_outputs(payload, Path(directory))
            saved = json.loads(paths["json"].read_text(encoding="utf-8"))
            with paths["csv"].open(newline="", encoding="utf-8") as handle:
                rows = list(csv.DictReader(handle))
        self.assertEqual(len(saved["trials"]), 36)
        self.assertEqual(len(rows), 36)
        self.assertEqual(
            {row["trial_key"] for row in rows},
            {trial["trial_key"] for trial in saved["trials"]},
        )


class PresentationOutputTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        inputs = study.load_study_inputs(ROOT)
        cls.payload = study.run_study(inputs, fake_trial)

    def test_matplotlib_is_headless(self):
        self.assertEqual(study.matplotlib.get_backend().lower(), "agg")

    def test_create_charts_writes_three_nonempty_png_files(self):
        with tempfile.TemporaryDirectory() as directory:
            paths = study.create_charts(
                self.payload["trials"],
                Path(directory),
            )
            self.assertEqual(
                set(paths),
                {"ranked_quality", "budget_convergence", "runtime"},
            )
            for path in paths.values():
                self.assertEqual(path.suffix, ".png")
                self.assertGreater(path.stat().st_size, 1_000)

    def test_report_has_required_sections_charts_and_page_breaks(self):
        with tempfile.TemporaryDirectory() as directory:
            folder = Path(directory)
            charts = study.create_charts(self.payload["trials"], folder)
            report = folder / "report.docx"
            study.build_report(self.payload, charts, report)
            document = Document(report)
            headings = [
                paragraph.text
                for paragraph in document.paragraphs
                if paragraph.style.name.startswith("Heading")
            ]
            self.assertEqual(
                headings,
                [
                    "Purpose and recommendation",
                    "Method and datasets",
                    "Parameter choices",
                    "Original dataset",
                    "Moderate stagger",
                    "Wide stagger",
                    "Cross-dataset comparison",
                    "Runtime and convergence",
                    "Limitations",
                    "Practical recommendation and reproduction",
                ],
            )
            self.assertEqual(len(document.inline_shapes), 3)
            xml = document._element.xml
            self.assertEqual(xml.count('w:type="page"'), 9)
            text = "\n".join(p.text.lower() for p in document.paragraphs)
            self.assertIn("one observed run", text)
            self.assertNotIn("statistically significant", text)
            self.assertNotIn("confidence interval", text)


if __name__ == "__main__":
    unittest.main()
