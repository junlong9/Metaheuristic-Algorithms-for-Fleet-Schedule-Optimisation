"""Build summary2.xlsx from the step-size study results, in the layout of summary1.xlsx.

Run from the repository root:
    python .scratch/vn-step-size-study/build_workbook.py
"""

from __future__ import annotations

from pathlib import Path
import json
import sys

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

sys.path.insert(0, str(Path(__file__).resolve().parent))
from study_data import (  # noqa: E402
    DATASET_LABELS,
    DATASETS,
    REPORT_DIR,
    format_steps,
    load,
    stop_text,
    trial,
)

OUTPUT = REPORT_DIR / "summary2.xlsx"
FONT = "Arial"
NAVY = "FF243B53"
INK = "FF172B4D"
BAND = "FFEDF2F7"
TRIALS_SHEET = "All trials"
HEADER_ROW = 7
FIRST_DATA_ROW = 8

COLUMNS = (
    ("Dataset", 23, None),
    ("Strategy", 36, None),
    ("Iteration limit", 14, "0"),
    ("Rank in dataset", 13, "0"),
    ("Minimum available", 16, "0"),
    ("Variability (SD)", 15, "0.0000"),
    ("Average available", 17, "0.0000"),
    ("Search passes", 14, "0"),
    ("Accepted moves", 15, "0"),
    ("Search stopped", 24, None),
    ("Runtime (sec)", 14, "0.00"),
    ("Rules met", 11, None),
    ("Ordered shift sets (days)", 72, None),
    ("Trial ID", 40, None),
    ("Strategy key", 22, None),
    ("Exact quality score", 50, None),
    ("No improving move", 18, None),
    ("Endpoint SHA256", 72, None),
)


def text_font(*, bold: bool = False, size: float = 11, colour: str = INK) -> Font:
    return Font(name=FONT, bold=bold, size=size, color=colour)


def put(ws, coordinate: str, value, **font) -> None:
    cell = ws[coordinate]
    cell.value = value
    cell.font = text_font(**font)
    cell.alignment = Alignment(vertical="center")


def header(ws, row: int, labels: list[str]) -> None:
    for column, label in enumerate(labels, start=1):
        cell = ws.cell(row, column, label)
        cell.font = text_font(bold=True, colour="FFFFFFFF")
        cell.fill = PatternFill("solid", fgColor=NAVY)
        cell.alignment = Alignment(
            horizontal="center", vertical="center", wrap_text=True
        )
    ws.row_dimensions[row].height = 40.05


def sort_key(item: dict, order: list[str], sets: dict) -> tuple:
    # Mirrors choose_recommendations so the first row of each dataset is its winner.
    return (
        DATASETS.index(item["dataset"]),
        item["dataset_rank"],
        item["max_iters"],
        sum(len(level) for level in sets[item["strategy_key"]]),
        item["elapsed_seconds"],
        order.index(item["strategy_key"]),
    )


def build_trials_sheet(ws, payload: dict) -> dict[str, int]:
    ws.sheet_view.showGridLines = False
    n_grid = sum(t["max_iters"] in payload["budgets"] for t in payload["trials"])
    n_ext = len(payload["extension_trials"])
    put(ws, "A2", "All optimiser trials", bold=True, size=16, colour=NAVY)
    put(
        ws,
        "A3",
        "Use the column filters to compare a dataset, strategy or iteration "
        "limit. Rank 1 is best; exact ties share a rank.",
    )
    put(
        ws,
        "A4",
        "Compare minimum availability first, then lower SD, then higher "
        "average. SD and availability are measured in aircraft.",
    )
    put(
        ws,
        "A5",
        f"Source: 2026-10-22 step-size study summary.csv. All "
        f"{len(payload['trials'])} trials retained ({n_grid} at 200/400/800 "
        f"plus {n_ext} at {payload['extension_budget']}). Technical "
        f"identifiers and exact scores are in columns M–R.",
    )
    header(ws, HEADER_ROW, [name for name, _, _ in COLUMNS])
    for index, (_, width, _) in enumerate(COLUMNS, start=1):
        ws.column_dimensions[ws.cell(1, index).column_letter].width = width

    order = payload["order"]
    sets = payload["sets"]
    rows = sorted(payload["trials"], key=lambda item: sort_key(item, order, sets))
    first_rows: dict[str, int] = {}
    for offset, item in enumerate(rows):
        row = FIRST_DATA_ROW + offset
        first_rows.setdefault(item["dataset"], row)
        values = (
            DATASET_LABELS[item["dataset"]],
            payload["labels"][item["strategy_key"]],
            item["max_iters"],
            item["dataset_rank"],
            item["min_availability"],
            item["availability_std"],
            item["mean_availability"],
            item["outer_passes"],
            item["accepted_moves"],
            stop_text(item),
            item["elapsed_seconds"],
            "Yes" if item["rule_compliant"] else "No",
            json.dumps(item["shift_sets"], separators=(",", ":")),
            item["trial_key"],
            item["strategy_key"],
            json.dumps(item["lexicographic_score"], separators=(",", ":")),
            item["natural_exhaustion"],
            item["endpoint_sha256"],
        )
        for column, (value, (_, _, number_format)) in enumerate(
            zip(values, COLUMNS), start=1
        ):
            cell = ws.cell(row, column, value)
            cell.font = text_font()
            cell.alignment = Alignment(
                vertical="center",
                horizontal="right" if number_format else None,
            )
            if number_format:
                cell.number_format = number_format
        ws.row_dimensions[row].height = 24.0
    for row in range(1, HEADER_ROW):
        ws.row_dimensions[row].height = 30.0 if row == 2 else 24.0
    last = FIRST_DATA_ROW + len(rows) - 1
    ws.freeze_panes = "C8"
    ws.auto_filter.ref = f"A{HEADER_ROW}:{ws.cell(HEADER_ROW, len(COLUMNS)).column_letter}{last}"
    winners = payload["recommendations"]["dataset_winners"]
    for dataset, row in first_rows.items():
        if rows[row - FIRST_DATA_ROW]["trial_key"] != winners[dataset]:
            raise AssertionError(f"{dataset}: first sorted row is not the winner")
    return first_rows


def build_summary_sheet(ws, payload: dict, first_rows: dict[str, int]) -> None:
    ws.sheet_view.showGridLines = False
    shared = payload["recommendations"]["shared"]
    key, budget = shared["strategy_key"], shared["max_iters"]
    if (key, budget) != ("live_wide_to_fine", 400):
        raise AssertionError("Summary text assumes the wide-to-fine ladder at 400")
    rec = {d: trial(payload, d, key, budget) for d in DATASETS}
    base = {d: trial(payload, d, "pairs_28", budget) for d in DATASETS}
    gaps = [base[d]["availability_std"] - rec[d]["availability_std"] for d in DATASETS]
    if not all(
        base[d]["min_availability"] == rec[d]["min_availability"] for d in DATASETS
    ) or min(gaps) <= 0:
        raise AssertionError("Baseline comparison text no longer holds")
    pairs = [(d, k) for d in DATASETS for k in payload["order"]]
    same = sum(
        trial(payload, d, k, 400)["endpoint_sha256"]
        == trial(payload, d, k, 800)["endpoint_sha256"]
        for d, k in pairs
    )
    capped_200 = sum(
        trial(payload, d, k, 200)["stop_reason"] == "iteration_cap" for d, k in pairs
    )
    moves = [rec[d]["accepted_moves"] for d in DATASETS]
    n_trials = len(payload["trials"])
    best_ext_rank = min(
        payload["by_key"][key]["dataset_rank"] for key in payload["extension_trials"]
    )
    if {key.split(":")[1] for key in payload["extension_trials"]} != {
        "pairs_1",
        "pairs_28_fine_first",
    }:
        raise AssertionError("Extension text assumes ±1 only and fine-first")
    p42 = {d: trial(payload, d, "pairs_42", budget) for d in DATASETS}
    helped_42 = sum(
        p42[d]["availability_std"] < base[d]["availability_std"] for d in DATASETS
    )

    put(ws, "A2", "Maintenance optimiser step-size results", bold=True, size=16, colour=NAVY)
    put(
        ws,
        "A4",
        f"Recommended shared setting: {payload['labels'][key]} with a "
        f"{budget}-iteration limit",
        bold=True,
    )
    put(ws, "A5", "Moves tested: ±56, ±42, ±28, ±14, ±7, ±3 and ±1 days, in that order.")
    put(ws, "A7", "Best tested setting for each dataset", bold=True)
    header(
        ws,
        8,
        [
            "Dataset",
            "Best strategy",
            "Iteration limit",
            "Minimum available",
            "Variability (SD)",
            "Average available",
            "Runtime (sec)",
        ],
    )
    references = (("A", None), ("B", None), ("C", "0"), ("E", "0"), ("F", "0.0000"), ("G", "0.0000"), ("K", "0.00"))
    for offset, dataset in enumerate(DATASETS):
        row = 9 + offset
        source = first_rows[dataset]
        for column, (letter, number_format) in enumerate(references, start=1):
            cell = ws.cell(row, column, f"='{TRIALS_SHEET}'!{letter}{source}")
            cell.font = text_font()
            cell.alignment = Alignment(vertical="center")
            if number_format:
                cell.number_format = number_format
            if offset % 2 == 0:
                cell.fill = PatternFill("solid", fgColor=BAND)

    row = 13
    put(ws, f"A{row}", f"Why use a {budget}-iteration limit", bold=True)
    for line in (
        f"400 and 800 reached exactly the same schedule in {same} of "
        f"{len(pairs)} dataset–strategy comparisons; every one that improved "
        f"at 800 had hit the 400 limit.",
        f"The wide-to-fine ladder finished naturally on all three datasets "
        f"after {min(moves)}–{max(moves)} accepted moves; 800 gave it no "
        f"quality gain.",
        f"200 reached its limit in all {capped_200} runs. Only the ±1-only and "
        f"fine-first ladders were still capped at 800; at 1600 they improved "
        f"but their best rank was {best_ext_rank} of "
        f"{len(payload['trials']) // len(DATASETS)}.",
        f"All {n_trials} schedules passed the implemented maintenance rules.",
    ):
        row += 1
        put(ws, f"A{row}", line)

    row += 2
    put(ws, f"A{row}", "Step-size findings", bold=True)
    for line in (
        f"Baseline (±28 to ±1) matched the recommended minimum on all three "
        f"datasets; its SD was higher by {min(gaps):.4f} to {max(gaps):.4f}.",
        "Each step removed from the top (±14, ±7, ±3, ±1) raised SD; ±3 and "
        "±1 also lost aircraft at the minimum.",
        "Order matters: the baseline steps searched fine step first gave the "
        "same schedules as ±1 only up to 800 iterations.",
        f"One combined set of the baseline steps was mixed and lost an "
        f"aircraft on Wide stagger; adding ±42 alone lowered SD on "
        f"{helped_42} of 3 datasets.",
    ):
        row += 1
        put(ws, f"A{row}", line)

    row += 2
    put(ws, f"A{row}", "Shift strategies", bold=True)
    for strategy in payload["order"]:
        row += 1
        put(ws, f"A{row}", payload["labels"][strategy], bold=True)
        put(ws, f"C{row}", format_steps(payload["sets"][strategy]))
    row += 1
    put(
        ws,
        f"A{row}",
        "Shifts are in days. Semicolons separate successive neighbourhoods; "
        "values in brackets are tried together; ± means both earlier and later.",
    )
    row += 1
    put(
        ws,
        f"A{row}",
        "The baseline reflects the user's recollection of the pre-study "
        "setting. The wide-to-fine ladder was adopted after the first study.",
    )

    row += 2
    for line in (
        "Search stopped: “No improving move” means none of the configured "
        "shifts improved the schedule.",
        "“Iteration limit reached” means the search used its full budget and "
        "may still improve with more passes.",
        "Runtime is one observed search run per trial, measured with four "
        "trials running in parallel; loading and validation are excluded.",
        "Reporting window: 21 January 2026 to 18 January 2038. Source: "
        "summary.csv and vn-hyperparameter-analysis2.docx in this folder.",
    ):
        put(ws, f"A{row}", line)
        row += 1

    for letter, width in zip("ABCDEFG", (23, 36, 13, 16, 15, 17, 14)):
        ws.column_dimensions[letter].width = width
    for index in range(1, row + 1):
        if index != 8:
            ws.row_dimensions[index].height = 30.0 if index == 2 else 24.0


def main() -> int:
    payload = load()
    workbook = Workbook()
    summary = workbook.active
    summary.title = "Summary"
    trials = workbook.create_sheet(TRIALS_SHEET)
    first_rows = build_trials_sheet(trials, payload)
    build_summary_sheet(summary, payload, first_rows)
    for sheet in (summary, trials):
        sheet.page_setup.orientation = "landscape"
        sheet.page_setup.paperSize = sheet.PAPERSIZE_A4
        sheet.sheet_properties.pageSetUpPr.fitToPage = True
        sheet.page_setup.fitToWidth = 1
        sheet.page_setup.fitToHeight = 0
    workbook.calculation.fullCalcOnLoad = True
    workbook.save(OUTPUT)
    print(f"Wrote {OUTPUT}; best rows {first_rows}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
