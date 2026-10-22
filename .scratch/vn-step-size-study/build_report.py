"""Build vn-hyperparameter-analysis2.docx from the step-size study results.

Run from the repository root:
    python .scratch/vn-step-size-study/build_report.py
"""

from __future__ import annotations

from collections import Counter
from pathlib import Path
import json
import sys

import matplotlib

matplotlib.use("Agg")
import matplotlib.pyplot as plt
from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

sys.path.insert(0, str(Path(__file__).resolve().parent))
from study_data import (  # noqa: E402
    DATASET_LABELS,
    DATASETS,
    ORIGINAL_KEYS,
    REPORT_DIR,
    SHA_CHECK,
    SHORT,
    QUESTION,
    format_steps,
    load,
    shared_table,
    trial,
)

OUTPUT = REPORT_DIR / "vn-hyperparameter-analysis2.docx"
FIGURES = Path(__file__).with_name("figures")
BUDGET = 400
FONT = "Arial"


def f4(value: float) -> str:
    return f"{value:.4f}"


def f2(value: float) -> str:
    return f"{value:.2f}"


def check(condition: bool, claim: str) -> None:
    if not condition:
        raise AssertionError(f"Report claim no longer holds: {claim}")


# ---------------------------------------------------------------- figures


def sd_chart(payload: dict, path: Path) -> None:
    rows = shared_table(payload, BUDGET)
    keys = [row["key"] for row in rows]
    figure, axes = plt.subplots(1, 3, figsize=(7.2, 3.5), sharey=True)
    for axis, dataset in zip(axes, DATASETS):
        cells = [trial(payload, dataset, key, BUDGET) for key in keys]
        best_min = max(item["min_availability"] for item in cells)
        positions = range(len(cells))
        for position, item in zip(positions, cells):
            best = item["min_availability"] == best_min
            axis.scatter(
                item["availability_std"],
                position,
                s=24,
                facecolor="black" if best else "white",
                edgecolor="black",
                linewidths=0.8,
                zorder=3,
            )
            axis.annotate(
                str(item["min_availability"]),
                (item["availability_std"], position),
                xytext=(5, 0),
                textcoords="offset points",
                va="center",
                fontsize=6.5,
            )
        low = min(item["availability_std"] for item in cells)
        high = max(item["availability_std"] for item in cells)
        axis.set_xlim(low * 0.92, high * 1.08)
        axis.set_title(DATASET_LABELS[dataset], fontsize=9)
        axis.tick_params(axis="x", labelsize=7)
        axis.grid(alpha=0.3)
    axes[0].set_yticks(range(len(keys)), [SHORT[key] for key in keys], fontsize=7)
    axes[0].invert_yaxis()
    figure.supxlabel(
        "Variability (SD) at 400 iterations, lower is better.\nNumber = minimum "
        "available aircraft; filled dots reach the best minimum for that forecast.",
        fontsize=7.5,
    )
    figure.tight_layout()
    figure.savefig(path, dpi=200)
    plt.close(figure)


def runtime_chart(payload: dict, path: Path) -> None:
    highlighted = {
        "live_wide_to_fine": ("*", 90, "black", "Wide-to-fine ladder (±56 to ±1)"),
        "pairs_28": ("s", 34, "white", "Baseline (±28 to ±1)"),
        "pairs_1": ("D", 26, "#808080", "±1 only"),
    }
    figure, axes = plt.subplots(1, 3, figsize=(7.2, 3.0))
    for axis, dataset in zip(axes, DATASETS):
        cells = [trial(payload, dataset, key, BUDGET) for key in payload["order"]]
        best_min = max(item["min_availability"] for item in cells)
        for item in cells:
            if item["strategy_key"] in highlighted:
                continue
            best = item["min_availability"] == best_min
            axis.scatter(
                item["elapsed_seconds"],
                item["availability_std"],
                s=18,
                marker="o" if best else "x",
                color="#4D4D4D" if best else "#808080",
                linewidths=1,
                label=(
                    "Other, best minimum" if best else "Other, lower minimum"
                ),
            )
        for key, (marker, size, fill, label) in highlighted.items():
            item = trial(payload, dataset, key, BUDGET)
            axis.scatter(
                item["elapsed_seconds"],
                item["availability_std"],
                s=size,
                marker=marker,
                facecolor=fill,
                edgecolor="black",
                linewidths=0.8,
                zorder=3,
                label=label,
            )
        axis.set_title(DATASET_LABELS[dataset], fontsize=9)
        axis.tick_params(labelsize=7)
        axis.grid(alpha=0.3)
    axes[0].set_ylabel("Variability (SD)", fontsize=7.5)
    handles = {}
    for axis in axes:
        for handle, label in zip(*axis.get_legend_handles_labels()):
            handles.setdefault(label, handle)
    figure.legend(
        handles.values(),
        handles.keys(),
        loc="lower center",
        ncols=3,
        fontsize=6.5,
        frameon=False,
    )
    figure.supxlabel(
        "Search time at 400 iterations (seconds, one run)",
        fontsize=7.5,
        y=0.17,
    )
    figure.tight_layout(rect=(0, 0.14, 1, 1))
    figure.savefig(path, dpi=200)
    plt.close(figure)


# ---------------------------------------------------------------- docx helpers


def set_cell_shading(cell, hex_fill: str) -> None:
    properties = cell._tc.get_or_add_tcPr()
    shading = OxmlElement("w:shd")
    shading.set(qn("w:val"), "clear")
    shading.set(qn("w:color"), "auto")
    shading.set(qn("w:fill"), hex_fill)
    properties.append(shading)


def style_document(document: Document) -> None:
    section = document.sections[0]
    section.page_width = Cm(21.0)
    section.page_height = Cm(29.7)
    for side in ("left_margin", "right_margin"):
        setattr(section, side, Cm(2.0))
    section.top_margin = Cm(1.8)
    section.bottom_margin = Cm(1.8)
    styles = document.styles
    for name, size, bold in (
        ("Normal", 10.5, False),
        ("Title", 20, True),
        ("Heading 1", 13, True),
    ):
        style = styles[name]
        style.font.name = FONT
        style.font.size = Pt(size)
        style.font.bold = bold
        style.font.color.rgb = RGBColor(0, 0, 0)
        fonts = style.element.rPr.rFonts
        for attribute in ("asciiTheme", "hAnsiTheme", "eastAsiaTheme", "cstheme"):
            fonts.attrib.pop(qn(f"w:{attribute}"), None)
        for attribute in ("ascii", "hAnsi", "eastAsia", "cs"):
            fonts.set(qn(f"w:{attribute}"), FONT)
    styles["Normal"].paragraph_format.space_after = Pt(5)
    styles["Normal"].paragraph_format.line_spacing = 1.08
    styles["Heading 1"].paragraph_format.space_before = Pt(10)
    styles["Heading 1"].paragraph_format.space_after = Pt(4)
    styles["Title"].paragraph_format.space_after = Pt(6)
    title_ppr = styles["Title"].element.get_or_add_pPr()
    for border in title_ppr.findall(qn("w:pBdr")):
        title_ppr.remove(border)


def add_table(
    document: Document,
    headers: list[str],
    rows: list[list[str]],
    widths: list[float],
    *,
    size: float = 8.5,
    bold_rows: set[int] = frozenset(),
) -> None:
    table = document.add_table(rows=1 + len(rows), cols=len(headers))
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    for column, width in zip(table.columns, widths):
        column.width = Cm(width)
    for row_index, values in enumerate([headers] + rows):
        for column, value in enumerate(values):
            cell = table.cell(row_index, column)
            cell.width = Cm(widths[column])
            paragraph = cell.paragraphs[0]
            paragraph.paragraph_format.space_after = Pt(0)
            paragraph.paragraph_format.line_spacing = 1.0
            if column > 0 and row_index > 0 and value[:1].isdigit():
                paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
            run = paragraph.add_run(value)
            run.font.size = Pt(size)
            run.bold = row_index == 0 or (row_index - 1) in bold_rows
            if row_index == 0:
                set_cell_shading(cell, "E6E6E6")
    spacer = document.add_paragraph()
    spacer.paragraph_format.space_after = Pt(0)
    spacer.paragraph_format.line_spacing = Pt(4)


def lead_paragraph(document: Document, lead: str, text: str) -> None:
    paragraph = document.add_paragraph()
    paragraph.add_run(f"{lead} ").bold = True
    paragraph.add_run(text)


def caption(document: Document, text: str) -> None:
    paragraph = document.add_paragraph()
    run = paragraph.add_run(text)
    run.italic = True
    run.font.size = Pt(8.5)


# ---------------------------------------------------------------- content


def build(payload: dict) -> Document:
    labels = payload["labels"]
    report_labels = dict(labels)
    report_labels.update(
        {
            "half_step_plus_42": "Half-step sets plus ±42",
            "live_wide_to_fine": "Wide to fine (±56 to ±1)",
        }
    )
    shared = payload["recommendations"]["shared"]
    winners = payload["recommendations"]["dataset_winners"]
    rec_key, rec_budget = shared["strategy_key"], shared["max_iters"]
    check(
        (rec_key, rec_budget) == ("live_wide_to_fine", BUDGET),
        "shared recommendation is the wide-to-fine ladder at 400",
    )

    rec = {dataset: trial(payload, dataset, rec_key, BUDGET) for dataset in DATASETS}
    baseline = {
        dataset: trial(payload, dataset, "pairs_28", BUDGET)
        for dataset in DATASETS
    }
    check(
        all(
            baseline[dataset]["min_availability"]
            == rec[dataset]["min_availability"]
            for dataset in DATASETS
        ),
        "baseline and recommendation have equal minimum availability",
    )
    check(
        all(
            baseline[dataset]["availability_std"]
            > rec[dataset]["availability_std"]
            for dataset in DATASETS
        ),
        "recommendation has lower variability than baseline",
    )
    sd_gap = {
        dataset: (
            baseline[dataset]["availability_std"]
            - rec[dataset]["availability_std"]
        )
        for dataset in DATASETS
    }

    document = Document()
    style_document(document)
    document.core_properties.title = "Variable-neighbourhood parameter analysis"
    document.core_properties.author = "UQ / Northrop Grumman fleet schedule project"
    document.add_heading("Variable-neighbourhood parameter analysis", level=0)

    document.add_paragraph(
        "Use the wide-to-fine shift ladder (±56, ±42, ±28, ±14, ±7, ±3 "
        "and ±1 days) with a 400-iteration limit. It gave the strongest "
        "overall result across the three forecasts and was best on the "
        "moderate and wide stagger forecasts. The baseline ladder (±28, ±14, "
        "±7, ±3 and ±1 days) reached the same minimum availability: 18, 19 "
        "and 20 aircraft, but its day-to-day availability varied slightly "
        "more. The difference in standard deviation was "
        f"{f4(min(sd_gap.values()))} to {f4(max(sd_gap.values()))} aircraft. "
        "The improvement is small, but it is consistent across all three "
        "forecasts."
    )

    document.add_heading("What we tested", level=1)
    document.add_paragraph(
        "The analysis compared 13 ordered shift strategies on the original, "
        "moderate stagger and wide stagger forecasts. Each strategy was tested "
        "at 200, 400 and 800 iterations. The two slowest-converging strategies "
        "were also tested at 1600 iterations, giving 123 trials in total. A "
        "shift moves a candidate maintenance event earlier or later, and the "
        "optimiser accepts the move that gives the best improvement."
    )
    document.add_paragraph(
        "Results were ranked by minimum available aircraft first, then lower "
        "standard deviation, then higher average availability. Standard "
        "deviation measures how much daily availability changes across the "
        "forecast. Availability was measured from 21 January 2026 to "
        "18 January 2038. All schedules passed the implemented maintenance "
        "checks."
    )

    document.add_heading("Best results", level=1)
    rows = []
    for dataset in DATASETS:
        item = payload["by_key"][winners[dataset]]
        rows.append(
            [
                DATASET_LABELS[dataset],
                report_labels[item["strategy_key"]],
                str(item["max_iters"]),
                str(item["min_availability"]),
                f4(item["availability_std"]),
                f4(item["mean_availability"]),
            ]
        )
    add_table(
        document,
        ["Forecast", "Best strategy", "Iterations", "Minimum", "SD", "Average"],
        rows,
        [3.0, 6.0, 1.8, 1.7, 1.7, 2.0],
    )
    document.add_paragraph(
        "The original forecast favoured the half-step ladder with ±42, while "
        "the wide-to-fine ladder was best on the two staggered forecasts."
    )

    table_400 = shared_table(payload, BUDGET)
    check(
        table_400[0]["key"] == rec_key,
        "recommended ladder ranks first across the three forecasts",
    )
    document.add_page_break()
    document.add_heading(f"All strategies at {BUDGET} iterations", level=1)
    comparison_rows = []
    for row in table_400:
        values = [report_labels[row["key"]]]
        for item in row["cells"]:
            capped = "*" if item["stop_reason"] == "iteration_cap" else ""
            values.append(
                f"{item['min_availability']} / "
                f"{f4(item['availability_std'])}{capped}"
            )
        comparison_rows.append(values)
    add_table(
        document,
        ["Strategy", "Original", "Moderate", "Wide"],
        comparison_rows,
        [7.0, 3.2, 3.2, 3.2],
        size=7.5,
        bold_rows={0},
    )
    document.add_paragraph(
        "Each result is minimum available aircraft / standard deviation. "
        "Lower standard deviation is better when minimum availability is "
        "equal. * indicates that the search reached the 400-iteration limit "
        "and may improve with more passes."
    )
    FIGURES.mkdir(exist_ok=True)
    sd_path = FIGURES / "sd-by-strategy-400.png"
    sd_chart(payload, sd_path)
    document.add_picture(str(sd_path), width=Cm(14.5))
    document.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
    caption(
        document,
        "Figure 1. Standard deviation at 400 iterations. Point labels show "
        "minimum available aircraft; filled points reach the best minimum for "
        "that forecast.",
    )

    document.add_page_break()
    document.add_heading("What the results show", level=1)
    lead_paragraph(
        document,
        "Large shifts help.",
        "Removing the largest shifts progressively increased variability. "
        "Ladders beginning at ±28, ±14 and ±7 retained the same minimum "
        "availability at 400 iterations, but ±3 and ±1 lost aircraft on at "
        "least two forecasts. In the recommended ladder, 36% of accepted "
        "moves were ±42 or ±56, showing that the larger shifts were actively "
        "used rather than simply adding search cost.",
    )
    lead_paragraph(
        document,
        "Search coarse to fine.",
        "Reversing the baseline ladder so that ±1 was searched first produced "
        "the same schedules as using ±1 alone through 800 iterations. Small "
        "improvements kept the search in the first neighbourhood, so it never "
        "reached the larger shifts. At 400 iterations this reduced the minimum "
        "availability from 18/19/20 aircraft to 17/18/18. Larger shifts should "
        "therefore be searched first.",
    )
    lead_paragraph(
        document,
        "Use 400 iterations.",
        "Every 200-iteration run reached its limit, so 200 is too short. The "
        "recommended ladder found no further improving move after 291 to 367 "
        "accepted moves on the three forecasts, and its 800-iteration results "
        "were identical. A 400-iteration limit captures its full result "
        "without unnecessary work.",
    )

    rt_path = FIGURES / "runtime-vs-sd-400.png"
    runtime_chart(payload, rt_path)
    document.add_picture(str(rt_path), width=Cm(14.5))
    document.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
    caption(
        document,
        "Figure 2. Search time against standard deviation at 400 iterations. "
        "Timings are single observations and are useful for relative "
        "comparison only.",
    )

    document.add_heading("Limits and next step", level=1)
    document.add_paragraph(
        "The analysis covers three related forecasts and deterministic runs, "
        "so it does not show variation across repeated trials or establish a "
        "global optimum. Differences between the leading strategies are only "
        "hundredths of an aircraft in standard deviation and may change on a "
        "different fleet or starting schedule. The optimiser also does not "
        "enforce the separate loose, medium and tight operational scenarios; "
        "those constraints still need review before deployment."
    )
    document.add_paragraph(
        "Keep the wide-to-fine ladder at 400 iterations as the shared setting. "
        "If simplicity is more important than the small reduction in "
        "variability, the ±28 baseline is a reasonable alternative because it "
        "preserves the same minimum availability. Repeat the comparison when "
        "the fleet, maintenance rules or starting schedule changes. The "
        "workbook contains the complete trial results."
    )
    return document


def _build_detailed_report(payload: dict) -> Document:
    sha = json.loads(SHA_CHECK.read_text(encoding="utf-8"))
    labels = payload["labels"]
    shared = payload["recommendations"]["shared"]
    winners = payload["recommendations"]["dataset_winners"]
    rec_key, rec_budget = shared["strategy_key"], shared["max_iters"]
    check(
        (rec_key, rec_budget) == ("live_wide_to_fine", BUDGET),
        "shared recommendation is the wide-to-fine ladder at 400",
    )
    table_400 = shared_table(payload, BUDGET)
    check(table_400[0]["key"] == rec_key, "recommended ladder ranks first at 400")
    rec = {d: trial(payload, d, rec_key, BUDGET) for d in DATASETS}
    base = {d: trial(payload, d, "pairs_28", BUDGET) for d in DATASETS}
    check(
        all(base[d]["min_availability"] == rec[d]["min_availability"] for d in DATASETS)
        and all(base[d]["availability_std"] > rec[d]["availability_std"] for d in DATASETS),
        "baseline matches minimum but has higher SD on every dataset",
    )
    sd_gap = {d: base[d]["availability_std"] - rec[d]["availability_std"] for d in DATASETS}
    base_rank = next(i for i, row in enumerate(table_400, 1) if row["key"] == "pairs_28")

    trials = payload["trials"]
    grid = [t for t in trials if t["max_iters"] in payload["budgets"]]
    extension = [payload["by_key"][key] for key in payload["extension_trials"]]
    n_strategies = len(payload["strategies"])
    n_grid = len(grid)
    n_total = len(trials)
    ranks_out_of = n_total // len(DATASETS)

    document = Document()
    style_document(document)
    document.core_properties.title = "Maintenance optimiser step-size results"
    document.core_properties.author = "UQ / Northrop Grumman fleet schedule project"

    document.add_heading("Maintenance optimiser step-size results", level=0)

    min_text = ", ".join(
        f"{rec[d]['min_availability']}" for d in DATASETS
    )
    document.add_paragraph(
        f"We recommend the wide-to-fine ladder (shifts of ±56, ±42, ±28, ±14, "
        f"±7, ±3 and ±1 days, searched in that order) with a {rec_budget}-"
        f"iteration limit as the shared setting for the three forecasts. Of the "
        f"{n_strategies} shift strategies tested it had the best combined "
        f"ranking, and it was the best tested setting on the moderate and wide "
        f"stagger forecasts. The baseline ladder (±28, ±14, ±7, ±3, ±1) comes "
        f"close: it reached the same minimum availability on all three "
        f"forecasts ({min_text} aircraft), but its daily availability varied "
        f"slightly more, with a standard deviation higher by "
        f"{f4(min(sd_gap.values()))} to {f4(max(sd_gap.values()))}. Ladders "
        f"that start below ±14, and the baseline searched fine step first, "
        f"were clearly worse."
    )
    document.add_paragraph(
        "This follow-up adds nine strategies to the four from the 20 October "
        "study so that the baseline and smaller step sizes can be compared "
        "directly. The baseline reflects the user's recollection of the setting "
        "used before the first study; it is not a recorded configuration and "
        "was not tested in that study. The wide-to-fine ladder was adopted "
        "after the first study."
    )

    # What we tested
    document.add_heading("What we tested", level=1)
    document.add_paragraph(
        f"Each strategy is an ordered list of day shifts. In every search pass "
        f"the optimiser tries the first set of shifts on the maintenance checks "
        f"that sit in the lowest-availability periods, accepts the best "
        f"improving move, and only moves to the next set when the current one "
        f"finds no improvement. We ran {n_strategies} strategies at limits of "
        f"{', '.join(str(b) for b in payload['budgets'][:-1])} and "
        f"{payload['budgets'][-1]} iterations on three starting forecasts, "
        f"giving {n_grid} trials. Two strategies were still improving when they "
        f"hit the 800 limit, so they were also run at "
        f"{payload['extension_budget']} iterations ({len(extension)} more "
        f"trials, {n_total} in total). The maintenance rules and selection "
        f"policy were unchanged, and availability was measured from "
        f"21 January 2026 to 18 January 2038."
    )
    document.add_paragraph(
        "Results were ranked by minimum available aircraft first, then lower "
        "standard deviation (SD, how much daily availability varies), then "
        "higher average availability. Exact quality ties favour the lower "
        "iteration limit and then the simpler strategy. "
        f"Rerunning the four original strategies reproduced all "
        f"{sha['compared'] - len(sha['mismatches'])} of {sha['compared']} final "
        f"schedules from the first study exactly (matching SHA256 fingerprints)."
    )
    rows = [
        [labels[key], format_steps(payload["sets"][key]), QUESTION[key]]
        for key in payload["order"]
    ]
    caption(document, "Table 1. Strategies and the shift steps they search, in order (days).")
    add_table(
        document,
        ["Strategy", "Shift steps, in search order", "Question"],
        rows,
        [5.6, 7.6, 3.8],
        size=8,
    )
    document.add_paragraph(
        "Semicolons separate sets that are searched one after another. Values "
        "in brackets are tried together as one set. ± means both earlier and "
        "later."
    ).runs[0].font.size = Pt(8.5)

    # Best results
    document.add_heading("Best results for each forecast", level=1)
    rows = []
    for dataset in DATASETS:
        item = payload["by_key"][winners[dataset]]
        rows.append([
            DATASET_LABELS[dataset],
            labels[item["strategy_key"]],
            str(item["max_iters"]),
            str(item["min_availability"]),
            f4(item["availability_std"]),
            f4(item["mean_availability"]),
        ])
    check(
        all(payload["by_key"][winners[d]]["strategy_key"] in ORIGINAL_KEYS for d in DATASETS),
        "no new strategy is a dataset winner",
    )
    caption(document, "Table 2. Best tested setting for each forecast.")
    add_table(
        document,
        ["Forecast", "Best strategy", "Iterations", "Minimum aircraft", "SD", "Average aircraft"],
        rows,
        [3.0, 5.8, 1.8, 2.0, 1.8, 2.2],
    )
    orig_rec = rec["original"]
    orig_best = payload["by_key"][winners["original"]]
    document.add_paragraph(
        f"None of the nine new strategies beat the first study's winners. On "
        f"the original forecast the overlapping half-step sets plus ±42 stayed "
        f"best; the recommended ladder ranked {orig_rec['dataset_rank']} of "
        f"{ranks_out_of} there, with the same minimum of "
        f"{orig_rec['min_availability']} aircraft and an SD only "
        f"{f4(orig_rec['availability_std'] - orig_best['availability_std'])} "
        f"higher. The forecasts start from different maintenance states, so "
        f"their availability levels should not be compared with each other or "
        f"read as gains caused by the optimiser."
    )

    # All strategies at 400
    document.add_heading(f"All strategies at {BUDGET} iterations", level=1)
    rows = []
    for row in table_400:
        values = [labels[row["key"]]]
        for item in row["cells"]:
            mark = "*" if item["stop_reason"] == "iteration_cap" else ""
            values.append(
                f"{item['min_availability']} / {f4(item['availability_std'])}{mark}"
            )
        values.append(str(row["rank_sum"]))
        rows.append(values)
    caption(
        document,
        f"Table 3. Minimum aircraft / SD at {BUDGET} iterations, sorted by "
        f"combined rank (sum of the three forecast ranks out of "
        f"{ranks_out_of}; lower is better). * = stopped at the iteration "
        f"limit and may improve with more passes.",
    )
    add_table(
        document,
        ["Strategy", "Original", "Moderate stagger", "Wide stagger", "Combined rank"],
        rows,
        [5.6, 2.8, 2.8, 2.8, 2.0],
        bold_rows={0},
    )
    FIGURES.mkdir(exist_ok=True)
    sd_path = FIGURES / "sd-by-strategy-400.png"
    sd_chart(payload, sd_path)
    document.add_picture(str(sd_path), width=Cm(15.5))
    document.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
    caption(
        document,
        f"Figure 1. Variability at {BUDGET} iterations for every strategy, in "
        f"Table 3 order. Note that each panel has its own scale.",
    )

    # Step-size findings
    document.add_heading("What the step-size experiments show", level=1)

    def at(key: str, budget: int = BUDGET) -> dict:
        return {d: trial(payload, d, key, budget) for d in DATASETS}

    def mins(cells: dict) -> str:
        return "/".join(str(cells[d]["min_availability"]) for d in DATASETS)

    def sds(cells: dict) -> str:
        return ", ".join(f4(cells[d]["availability_std"]) for d in DATASETS)

    ladder = ["pairs_28", "pairs_14", "pairs_7", "pairs_3", "pairs_1"]
    shrink = {key: at(key) for key in ladder}
    p1_800 = at("pairs_1", 800)
    p1_1600 = {d: trial(payload, d, "pairs_1", 1600) for d in DATASETS}
    check(mins(shrink["pairs_14"]) == mins(base), "±14 keeps the minimum")
    check(mins(shrink["pairs_7"]) == mins(base), "±7 keeps the minimum at 400")
    check(
        all(
            shrink[a][d]["availability_std"] < shrink[b][d]["availability_std"]
            for a, b in zip(ladder, ladder[1:])
            for d in DATASETS
        ),
        "SD rises with every step removed",
    )
    lead_paragraph(
        document,
        "Shrinking the largest step.",
        f"Removing the largest step one at a time "
        f"made variability worse at every stage on every forecast. Minimum "
        f"aircraft (original/moderate/wide) at {BUDGET} iterations were "
        f"{mins(base)} for the ±28 baseline, {mins(shrink['pairs_14'])} for "
        f"±14, {mins(shrink['pairs_7'])} for ±7, {mins(shrink['pairs_3'])} for "
        f"±3 and {mins(shrink['pairs_1'])} for ±1 alone. SD rose from "
        f"{sds(base)} with ±28 to {sds(shrink['pairs_14'])} with ±14 and "
        f"{sds(shrink['pairs_7'])} with ±7. The ±7, ±3 and ±1 ladders also ran "
        f"out of iterations at {BUDGET}. Even with more budget, ±1 alone fell "
        f"short of the baseline's {mins(base)}: it reached {mins(p1_800)} "
        f"aircraft at 800 and {mins(p1_1600)} at 1600. Steps of at least ±14 "
        f"are needed, and ±28 is clearly better "
        f"than ±14."
    )
    fine = at("pairs_28_fine_first")
    check(
        all(
            trial(payload, d, "pairs_28_fine_first", b)["endpoint_sha256"]
            == trial(payload, d, "pairs_1", b)["endpoint_sha256"]
            for d in DATASETS
            for b in payload["budgets"]
        ),
        "fine-first equals ±1 only up to 800",
    )
    check(
        all(
            abs(step["shift"]) == 1
            for d in DATASETS
            for b in payload["budgets"]
            for step in trial(payload, d, "pairs_28_fine_first", b)["trace"][1:]
        ),
        "fine-first accepts only ±1 moves up to 800",
    )
    fine_1600 = {d: trial(payload, d, "pairs_28_fine_first", 1600) for d in DATASETS}
    lead_paragraph(
        document,
        "Order matters.",
        f"Searching the same five baseline steps fine step "
        f"first (±1, ±3, ±7, ±14, ±28) produced exactly the same schedules as "
        f"±1 alone at 200, 400 and 800 iterations. Every accepted move in "
        f"those runs was ±1: a ±1 move kept finding some small improvement, "
        f"so the larger steps were never reached. At {BUDGET} it reached "
        f"{mins(fine)} aircraft against "
        f"{mins(base)} for the baseline order, and it was still at the limit "
        f"after 1600 iterations on every forecast ({mins(fine_1600)} aircraft)."
        f" Large steps must come first."
    )
    single = at("single_set_28")
    single_200 = at("single_set_28", 200)
    base_200 = at("pairs_28", 200)
    check(
        single["original"]["availability_std"] < base["original"]["availability_std"]
        and single["moderate"]["availability_std"] > base["moderate"]["availability_std"]
        and single["wide"]["min_availability"] < base["wide"]["min_availability"],
        "single set mixed: better original, worse moderate, loses wide minimum",
    )
    lead_paragraph(
        document,
        "Single set versus ladder.",
        f"Trying all ten baseline shifts together "
        f"as one set was slightly better than the ladder on the original "
        f"forecast (SD {f4(single['original']['availability_std'])} against "
        f"{f4(base['original']['availability_std'])}), worse on moderate "
        f"stagger ({f4(single['moderate']['availability_std'])} against "
        f"{f4(base['moderate']['availability_std'])}) and lost an aircraft at "
        f"the minimum on wide stagger ({single['wide']['min_availability']} "
        f"against {base['wide']['min_availability']}). Each pass is also "
        f"heavier: at 200 iterations it took "
        f"{f2(min(single_200[d]['elapsed_seconds'] for d in DATASETS))}–"
        f"{f2(max(single_200[d]['elapsed_seconds'] for d in DATASETS))} s "
        f"against {f2(min(base_200[d]['elapsed_seconds'] for d in DATASETS))}–"
        f"{f2(max(base_200[d]['elapsed_seconds'] for d in DATASETS))} s for "
        f"the ladder. The ladder is the better default."
    )
    p42 = at("pairs_42")
    check(
        p42["original"]["availability_std"] < base["original"]["availability_std"]
        and p42["moderate"]["availability_std"] < base["moderate"]["availability_std"]
        and p42["wide"]["availability_std"] > base["wide"]["availability_std"],
        "±42 helps original and moderate, hurts wide",
    )
    wide_moves = Counter()
    all_moves = 0
    for d in DATASETS:
        for step in rec[d]["trace"][1:]:
            wide_moves[abs(step["shift"]) >= 42] += 1
            all_moves += 1
    share = 100 * wide_moves[True] / all_moves
    mean_gap = max(
        abs(rec[d]["mean_availability"] - base[d]["mean_availability"])
        for d in DATASETS
    )
    lead_paragraph(
        document,
        "Wide steps.",
        f"Adding ±42 in front of the baseline improved SD on the "
        f"original and moderate forecasts ({f4(p42['original']['availability_std'])} "
        f"and {f4(p42['moderate']['availability_std'])}) but made wide stagger "
        f"worse ({f4(p42['wide']['availability_std'])} against "
        f"{f4(base['wide']['availability_std'])}). Adding both ±56 and ±42 "
        f"improved SD on all three forecasts, by "
        f"{', '.join(f4(sd_gap[d]) for d in DATASETS)}, with the same minimum "
        f"and average availability within {mean_gap:.3f} aircraft. The wide moves are "
        f"genuinely used: {share:.0f}% of the recommended ladder's accepted "
        f"moves were ±42 or ±56. The gain over the baseline is real but small."
    )
    o14 = at("overlap_14")
    half = at("historical_half_step")
    check(
        o14["wide"]["min_availability"] > half["wide"]["min_availability"]
        and o14["original"]["availability_std"] > half["original"]["availability_std"],
        "overlap without ±28 mixed",
    )
    lead_paragraph(
        document,
        "Overlapping sets.",
        f"Dropping ±28 from the overlapping half-step sets "
        f"raised the wide stagger minimum from {half['wide']['min_availability']} "
        f"to {o14['wide']['min_availability']} aircraft, but made SD worse on "
        f"the original forecast and, at {BUDGET} iterations, ran out of "
        f"iterations on moderate stagger. Neither overlapping variant beat the "
        f"simple paired ladders overall."
    )

    # Iteration limit
    document.add_heading(f"Why use {BUDGET} iterations", level=1)
    capped_200 = sum(
        trial(payload, d, key, 200)["stop_reason"] == "iteration_cap"
        for d in DATASETS for key in payload["order"]
    )
    pairs = [(d, key) for d in DATASETS for key in payload["order"]]
    same = [
        (d, key) for d, key in pairs
        if trial(payload, d, key, 400)["endpoint_sha256"]
        == trial(payload, d, key, 800)["endpoint_sha256"]
    ]
    capped_400 = [
        (d, key) for d, key in pairs
        if trial(payload, d, key, 400)["stop_reason"] == "iteration_cap"
    ]
    check(set(capped_400) == set(pairs) - set(same), "400/800 differences are capped runs")
    check(capped_200 == len(pairs), "every 200 run capped")
    moves = [rec[d]["accepted_moves"] for d in DATASETS]
    document.add_paragraph(
        f"All {capped_200} runs at 200 iterations hit their limit, so 200 is "
        f"too short. 400 and 800 produced identical schedules in "
        f"{len(same)} of {len(pairs)} strategy–forecast comparisons. Every one "
        f"of the {len(pairs) - len(same)} that improved at 800 had already "
        f"stopped at the 400 limit: the ±7, ±3, ±1 and fine-first ladders on "
        f"every forecast, plus dense pairs, ±14 and overlapping sets without "
        f"±28 on moderate stagger. The recommended ladder finished with no "
        f"improving move on all three forecasts after {min(moves)}–"
        f"{max(moves)} accepted moves, so 800 adds nothing for it."
    )
    ext_text = []
    for key in ("pairs_1", "pairs_28_fine_first"):
        rows_ext = [trial(payload, d, key, 1600) for d in DATASETS]
        finished = [DATASET_LABELS[r["dataset"]].lower() for r in rows_ext if r["stop_reason"] == "local_exhaustion"]
        ext_text.append((key, rows_ext, finished))
    p1_rows, p1_finished = ext_text[0][1], ext_text[0][2]
    ff_finished = ext_text[1][2]
    best_ext_rank = min(r["dataset_rank"] for _, rows_ext, _ in ext_text for r in rows_ext)
    document.add_paragraph(
        f"Only ±1 alone and the fine-first ladder were still at the limit after "
        f"800 iterations, so they were rerun at 1600. Doubling the budget "
        f"helped them (±1 alone improved from {mins(p1_800)} to "
        f"{mins({r['dataset']: r for r in p1_rows})} aircraft) and ±1 alone "
        f"finished on {' and '.join(p1_finished)}, but fine-first "
        f"{'was still at the limit on every forecast' if not ff_finished else 'finished on ' + ' and '.join(ff_finished)}. "
        f"Their best 1600-iteration rank was {best_ext_rank} of "
        f"{ranks_out_of}. More budget does not make up for missing large steps."
    )
    rec_times = [rec[d]["elapsed_seconds"] for d in DATASETS]
    base_times = [base[d]["elapsed_seconds"] for d in DATASETS]
    first = {}
    with (REPORT_DIR.parent / "2026-10-20-vn-hyperparameter-analysis/summary.csv").open(encoding="utf-8") as handle:
        import csv

        for row in csv.DictReader(handle):
            first[row["trial_key"]] = float(row["elapsed_seconds"])
    old_rec = [first[f"{d}:{rec_key}:{BUDGET}"] for d in DATASETS]
    times = [t["elapsed_seconds"] for t in trials]
    check(all(t["rule_compliant"] for t in trials), "all trials rule compliant")
    document.add_paragraph(
        f"All {n_total} schedules passed the implemented maintenance checks. "
        f"Search times ranged from {f2(min(times))} to {f2(max(times))} "
        f"seconds. The recommended setting took {f2(min(rec_times))}–"
        f"{f2(max(rec_times))} s per forecast and the baseline at {BUDGET} took "
        f"{f2(min(base_times))}–{f2(max(base_times))} s (Figure 2). These times "
        f"were measured with four trials running at once on a four-core cloud "
        f"machine, so they run higher than the first study's single runs "
        f"({f2(min(old_rec))}–{f2(max(old_rec))} s for the same setting). Each "
        f"is a single observation and excludes loading and validation; treat "
        f"them as relative, not absolute."
    )
    rt_path = FIGURES / "runtime-vs-sd-400.png"
    runtime_chart(payload, rt_path)
    document.add_picture(str(rt_path), width=Cm(15.0))
    document.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
    caption(
        document,
        f"Figure 2. Search time against variability at {BUDGET} iterations. "
        f"Faster settings are the small-step ladders, which give worse schedules. "
        f"The fine-first ladder sits under ±1 only because its schedules are identical.",
    )

    # Limits
    document.add_heading("Limits and next step", level=1)
    document.add_paragraph(
        f"These results cover three related forecasts and {n_strategies} "
        f"strategies. Each trial is a single deterministic run, so there is no "
        f"spread across repeats, and the timings are single observations. The "
        f"SD differences between the leading strategies are small (hundredths "
        f"of an aircraft) and could reorder on other forecasts. The results do "
        f"not establish a global optimum. The optimiser does not enforce the "
        f"separate loose, medium and tight operational scenarios, so those "
        f"constraints still need review before adoption. The baseline is based "
        f"on recollection rather than a recorded configuration."
    )
    document.add_paragraph(
        f"Keep the wide-to-fine ladder at {BUDGET} iterations. If a simpler "
        f"setting is preferred, the ±28 baseline at {BUDGET} keeps the same "
        f"minimum availability on these forecasts at a small cost in "
        f"variability (it ranked {base_rank} of {n_strategies} at {BUDGET}). "
        f"Do not use ladders that start below ±14 or that search fine steps "
        f"first. Rerun the comparison when the fleet, maintenance rules or "
        f"starting state changes."
    )
    return document


def main() -> int:
    payload = load()
    document = build(payload)
    document.save(OUTPUT)
    print(f"Wrote {OUTPUT}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
