"""Write the priority-VNS hyperparameter Word report."""

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt


OUTPUT = "hyperparameter-report.docx"


def add_heading(doc, text, level=1):
    doc.add_heading(text, level=level)


def add_para(doc, text):
    paragraph = doc.add_paragraph(text)
    paragraph.paragraph_format.space_after = Pt(8)
    return paragraph


def add_table(doc, headers, rows):
    table = doc.add_table(rows=1 + len(rows), cols=len(headers))
    table.style = "Table Grid"
    for index, header in enumerate(headers):
        table.rows[0].cells[index].text = header
    for row_index, row in enumerate(rows, start=1):
        for col_index, value in enumerate(row):
            table.rows[row_index].cells[col_index].text = str(value)
    doc.add_paragraph()


def main():
    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(11)

    title = doc.add_heading(
        "Priority-based optimiser hyperparameter report",
        0,
    )
    title.alignment = WD_ALIGN_PARAGRAPH.LEFT

    add_para(
        doc,
        "This report is the write-up of a 74-trial study of "
        "vn_priority_based_selection.py on the local "
        "heavy_maintenance_forecast.xlsx workbook. The production "
        "module was not changed. A quiet copy of the search loop "
        "reused the live scoring, rebuild, and neighbour code. "
        "The algorithm is deterministic, so each setting was run once. "
        "Every finished schedule passed is_rule_compliant_schedule.",
    )

    add_heading(doc, "Verdict", 1)
    add_para(
        doc,
        "Minimum availability 19 was not found. The floor stayed at 18. "
        "The published default still reaches min 18 with standard "
        "deviation 1.413971 in about 15 seconds, which matches the "
        "Iteration 1 hill-climb. Those neighbourhoods are not the best "
        "ones. The best setting found a smoother 18 at standard "
        "deviation 1.391860.",
    )
    add_para(
        doc,
        "That smoothness also beats the 1.405740 keep-gate used to "
        "judge the later annealing work. A better neighbourhood on the "
        "original hill-climber cleared the target that annealing was "
        "asked to meet, without leaving the deterministic search.",
    )

    add_heading(doc, "Recommended settings", 1)
    add_para(
        doc,
        "Keep first-priority-level selection and a 400-iteration cap. "
        "Change the neighbourhoods to a half-step VNS ladder:",
    )
    add_para(
        doc,
        "shift_sets = [(-28, -14, 14, 28), (-14, -7, 7, 14), "
        "(-7, -3, 3, 7), (-3, -1, 1, 3)]",
    )
    add_para(
        doc,
        "On this workbook that run accepted 319 moves, stopped because "
        "nothing else improved, and finished in 96.5 seconds. Mean "
        "availability stayed at 21.713010, as it did in every trial, "
        "because check durations are fixed.",
    )
    add_para(
        doc,
        "If you want ordinary plus-or-minus pairs only, the next best "
        "ladder is (-28, 28), (-21, 21), (-14, 14), (-7, 7), (-3, 3), "
        "(-1, 1) at std 1.396502. A cheap upgrade from the published "
        "default is to put (-42, 42) in front of it, which reaches "
        "1.404311 in 20 seconds.",
    )

    add_heading(doc, "How the study was run", 1)
    add_para(
        doc,
        "The local workbook has 24 aircraft, 278 events, and a 4,480-day "
        "horizon. Off the sheet the fleet is min 15, std 1.950. "
        "I first reproduced the published default, then swept iteration "
        "caps, single jump sizes, VNS ladders, neighbourhood order, "
        "richer offset tuples, and selection policy. After that I reran "
        "the six best distinct settings at 80, 200, and 800 iterations "
        "and again with all-levels selection. Total wall time was "
        "3,151 seconds. One repeated winner sat for 16 minutes on a "
        "busy machine; the first time that setting finished in 95 "
        "seconds. I treat the long repeat as machine noise.",
    )

    add_heading(doc, "Iteration budget", 1)
    add_para(
        doc,
        "On the default ladder the search accepts 280 moves and then "
        "stops. A cap of 400 or 800 gives the same schedule. Caps of "
        "160 and below never reach min 18. The later winners needed "
        "roughly 260 to 350 accepted moves as well. Do not cut "
        "max_iters to 200 and assume the floor will hold.",
    )
    add_table(
        doc,
        ["Cap", "Minimum", "Std", "Moves"],
        [
            ["5", "16", "1.873", "5"],
            ["10", "17", "1.810", "10"],
            ["20", "17", "1.758", "20"],
            ["40", "17", "1.675", "40"],
            ["80", "17", "1.569", "80"],
            ["160", "17", "1.468", "160"],
            ["400", "18", "1.414", "280"],
            ["800", "18", "1.414", "280"],
        ],
    )

    add_heading(doc, "Single neighbourhoods cannot replace VNS", 1)
    add_para(
        doc,
        "Only ±5, ±7, and ±14 reached min 18 on their own. The best of "
        "those, ±7, stopped at std 1.471, well behind the default "
        "ladder. ±10 and ±3 stalled at 17. ±1 and ±2 spent 400 tiny "
        "moves and also stayed at 17. ±56 and ±70 barely moved, because "
        "the rule rebuild refuses most of those jumps. Large steps lift "
        "the floor. Small steps polish a schedule that is already there. "
        "You need both, in that order.",
    )

    add_heading(doc, "Order", 1)
    add_para(
        doc,
        "Small-first versions of the default and weekly ladders never "
        "reached min 18. They spent the budget taking the same tiny "
        "improving moves that a ±3 or ±1 neighbourhood would take. An "
        "interleaved order still reached 18 and slightly beat the "
        "default. Starting at ±42 and immediately dropping to ±1 did "
        "not. Keep large jumps first.",
    )

    add_heading(doc, "Ladder shape", 1)
    add_para(
        doc,
        "A coarse ladder of only ±70, ±42, and ±28 stayed at min 17. "
        "Two-level shortcuts reached 18 but left a bumpier series. "
        "Adding a one-day last step or a 42-day first step both beat "
        "the published default. Filling the gaps between the default "
        "rungs was better still. The best ordinary ±k ladder was the "
        "dense weekly set ±28, ±21, ±14, ±7, ±3, ±1.",
    )

    add_heading(doc, "Richer offsets", 1)
    add_para(
        doc,
        "The winner was not a new jump size. It was extra offsets on "
        "each existing rung, so a 28-day level can also try 14 days, "
        "and a 3-day level can also try one day. Putting every offset "
        "into one giant union was worse than keeping the layers. The "
        "layering is the thing that makes VNS work. Asymmetric left- "
        "or right-heavy sets did not help.",
    )

    add_heading(doc, "Selection policy", 1)
    add_para(
        doc,
        "The live first-priority-level rule beat both alternatives. "
        "Looking across every priority level on the default ladder was "
        "slightly worse and more than four times slower. First-improving "
        "was worse again. Forcing all-levels selection onto the winning "
        "neighbourhoods also lost smoothness. Restricting attention to "
        "events on the worst days is aligned with the objective, not "
        "just a speed hack.",
    )

    add_heading(doc, "Best first-pass runs that held min 18", 1)
    add_table(
        doc,
        ["Setting", "Std", "Moves", "Seconds"],
        [
            ["Half-step VNS ladder (winner)", "1.391860", "319", "96.5"],
            ["Dense ±k ladder", "1.396502", "329", "82.7"],
            ["Geometric plus one day", "1.399695", "346", "94.9"],
            ["Wide dense ladder", "1.402084", "339", "28.9"],
            ["Paired fortnight / week", "1.403516", "261", "41.6"],
            ["Default plus ±42", "1.404311", "291", "19.5"],
            ["Geometric ±32..±2", "1.404470", "333", "22.4"],
            ["Published default", "1.413971", "280", "14.6"],
        ],
    )

    add_heading(doc, "What I would not do", 1)
    add_para(
        doc,
        "Do not start with small steps. Do not collapse the ladder into "
        "one union of every offset. Do not cut the iteration cap to 200 "
        "and assume min 18 will survive. Do not change the selection "
        "rule. And do not treat this study as proof that annealing "
        "cannot find min 19. It only says the old neighbourhoods were "
        "leaving smoothness on the table. A later escape experiment "
        "should warm-start from this winner, not from the published "
        "default.",
    )

    add_heading(doc, "Reproduce", 1)
    add_para(
        doc,
        "python .scratch/priority-based-selection/run_hyperparameter_study.py",
    )
    add_para(
        doc,
        "The script skips names already stored in hyperparameter-results.json.",
    )

    doc.save(OUTPUT)
    print(f"Wrote {OUTPUT}")


if __name__ == "__main__":
    main()
