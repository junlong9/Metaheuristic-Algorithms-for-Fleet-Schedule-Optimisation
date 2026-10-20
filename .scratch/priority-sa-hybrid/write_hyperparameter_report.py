"""Write the basin-hop hyperparameter Word report from measured scores."""

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt


OUTPUT = "hyperparameter-report.docx"


def add_heading(doc, text, level=1):
    doc.add_heading(text, level=level)


def add_para(doc, text):
    paragraph = doc.add_paragraph(text)
    paragraph.paragraph_format.space_after = Pt(8)
    return paragraph


def add_table(doc, headers, rows):
    table = doc.add_table(rows=1 + len(rows), cols=len(headers))
    table.style = "Table Grid"
    for index, header in enumerate(headers):
        table.rows[0].cells[index].text = header
    for row_index, row in enumerate(rows, start=1):
        for col_index, value in enumerate(row):
            table.rows[row_index].cells[col_index].text = str(value)
    doc.add_paragraph()


def main():
    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(11)

    title = doc.add_heading("Basin-hop hyperparameter report", 0)
    title.alignment = WD_ALIGN_PARAGRAPH.LEFT

    add_para(
        doc,
        "This report records which basin-hop knobs produced the best "
        "lexicographic score on the local heavy_maintenance_forecast.xlsx "
        "workbook, and which earlier study results do not apply to that file.",
    )

    add_heading(doc, "Verdict", 1)
    add_para(
        doc,
        "Minimum availability 19 was not found on the local workbook. "
        "The printed score from python savn_priority_sa.py on 17 September 2026 "
        "was min 18, std 1.406534, avg 21.713010 in 239.4 seconds.",
    )
    add_para(
        doc,
        "That result beats the frozen hill-climb std of 1.413971. "
        "It does not beat the earlier local basin-hop std of 1.405740. "
        "The spec keep gate for a smoother 18 is std below 1.405740. "
        "This tabu and plateau run did not meet that gate.",
    )
    add_para(
        doc,
        "The best measured configuration on the real workbook is still the "
        "earlier local run with 15 cycles, 8 Metropolis moves, temperature 1.5, "
        "400 polish iterations, and seed 99. That run printed min 18 and "
        "std 1.405740 in 95.9 seconds.",
    )

    add_heading(doc, "Success gate", 1)
    add_para(
        doc,
        "The score is unchanged. Higher is better in this order: minimum "
        "availability, then lower standard deviation, then higher average.",
    )
    add_table(
        doc,
        ["Printed outcome", "Decision"],
        [
            ["min 19 or higher", "Success. Keep the change."],
            ["min 18 and std below 1.405740", "Fallback keep."],
            ["min 18 and std 1.405740 or higher", "The change did not earn its keep."],
        ],
    )
    add_para(
        doc,
        "Claim 19 only from a fresh printed score on the existing local "
        "heavy_maintenance_forecast.xlsx file. Cloud scores and regenerated "
        "forecast scores do not count.",
    )

    add_heading(doc, "Scores that count on the local workbook", 1)
    add_table(
        doc,
        ["Run", "Min", "Std", "Avg", "Time"],
        [
            [
                "Frozen hill-climb in vn_priority_based_selection.py",
                "18",
                "1.413971",
                "21.713010",
                "Not re-timed here",
            ],
            [
                "Local basin-hop 15 cycles, 8 moves, T=1.5, seed 99",
                "18",
                "1.405740",
                "21.713010",
                "95.9 s",
            ],
            [
                "This run. Best-so-far start, 30 plateau moves, tabu polish, "
                "15 cycles, T=1.5, seed 99",
                "18",
                "1.406534",
                "21.713010",
                "239.4 s",
            ],
        ],
    )
    add_para(
        doc,
        "Warm start for this run finished at min 18 and std 1.413971 after "
        "280 polish iterations. That matches the published hill-climb. "
        "Escape cycles 1, 3, 5, and 10 then lowered std to 1.410337, "
        "1.410178, 1.407010, and 1.406534. No cycle printed min 19.",
    )

    add_heading(doc, "What the knobs did on the regenerated cloud file", 1)
    add_para(
        doc,
        "A 146-trial cloud grid ran for 3183.7 seconds on a regenerated "
        "forecast, not on the local workbook. Every trial returned min 16. "
        "Those numbers cannot answer the 18 to 19 question. They only show "
        "which knobs smoothed a min-16 basin.",
    )
    add_para(
        doc,
        "On that regenerated file the best tied std was 1.479924. The "
        "winning settings were 15 cycles, 8 escape moves, temperature 1.5 "
        "or 2.0, 400 polish iterations, and seed 99 or 123.",
    )
    add_table(
        doc,
        ["Knob", "What worked better", "Evidence"],
        [
            [
                "Temperature",
                "1.5 and 2.0 beat 0.6",
                "Mean std 1.488 at T=2.0, 1.489 at T=1.5, 1.494 at T=0.6. "
                "T=0.6 skipped many first hops.",
            ],
            [
                "Escape depth",
                "8 moves beat 1 move",
                "Mean std 1.485 versus 1.495. Deeper walks helped smoothness only.",
            ],
            [
                "Cycles",
                "15 beat 5",
                "The three tied best trials all used 15 cycles.",
            ],
            [
                "Polish iterations",
                "400 beat 200",
                "200 never finished the warm start and was worse in every trial.",
            ],
            [
                "Seed",
                "99 and 123 tied at the top",
                "seed99_T1.5_m8_c15_it400 and two seed-123 twins all hit 1.479924.",
            ],
        ],
    )
    add_para(
        doc,
        "Those knobs became the default experiment for the local spec. "
        "escape_moves was then replaced by one forced floor-drop plus "
        "plateau_moves=30. The local verification used that new loop, "
        "not another grid.",
    )

    add_heading(doc, "What this local verification showed", 1)
    add_para(
        doc,
        "Each of the 15 cycles accepted a floor-drop and completed the "
        "30-move plateau. Polish always walked the escaped plan back to "
        "min 18. After the forced drop the walk sat at min 16 for most "
        "plateau steps, then greedy_polish climbed 16 to 17 to 18.",
    )
    add_para(
        doc,
        "Tabu blocked the exact pre-escape start vector. It did not stop "
        "polish from reaching another min-18 plan. Starting the next cycle "
        "from the protected best kept later hops from inheriting a worse "
        "local opt. That is why std still improved after a worse polish.",
    )
    add_para(
        doc,
        "The extra plateau length cost time and did not buy a better floor. "
        "The earlier 8-move local run was faster and printed a slightly "
        "lower std.",
    )

    add_heading(doc, "What worked best", 1)
    add_para(
        doc,
        "On the workbook that matters, keep temperature 1.5, 15 cycles, "
        "400 polish iterations, and seed 99. Use a short Metropolis walk "
        "after the floor-drop. Eight moves beat both a one-move hop on the "
        "cloud file and this 30-move plateau on the local file.",
    )
    add_para(
        doc,
        "Do not treat the cloud min-16 ranking as a local ranking. The "
        "regenerated forecast warm-starts at min 16. The local file "
        "warm-starts at min 18.",
    )

    add_heading(doc, "What did not work", 1)
    add_para(
        doc,
        "Temperature 0.6 skipped first hops too often on the cloud grid. "
        "Polish capped at 200 iterations never finished the warm start. "
        "A 30-move plateau plus tabu did not find min 19 and did not beat "
        "std 1.405740.",
    )

    add_heading(doc, "How to reproduce the local score", 1)
    add_para(
        doc,
        "Run python savn_priority_sa.py against the existing "
        "heavy_maintenance_forecast.xlsx file. Do not regenerate that file. "
        "Read the printed Best score line. The 17 September 2026 run printed "
        "Best score: (18.0, -1.4065339515503619, 21.713010488730195) and "
        "Completed escape cycles: 15.",
    )

    doc.save(OUTPUT)
    print(f"Wrote {OUTPUT}")


if __name__ == "__main__":
    main()
