from pathlib import Path
import sys,json,math,statistics,hashlib
from collections import Counter
import numpy as np
from docx import Document
from docx.shared import Inches,Pt,RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT,WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from reportlab.graphics.shapes import Drawing,String,Line,Rect,PolyLine,Circle
from reportlab.graphics import renderPDF
import pypdfium2 as pdfium

ROOT=Path(__file__).resolve().parents[2]; WORK=Path(__file__).resolve().parent
sys.path.insert(0,str(WORK))
import run_analysis as analysis
opt=analysis.opt
hist=json.loads((ROOT/'.scratch/priority-based-selection/hyperparameter-results.json').read_text())
fresh=json.loads((WORK/'fresh-results.json').read_text())
assert fresh.get('complete') and fresh['production_equivalence_verified']
H=hist['trials']; F=fresh['trials']
by_h={t['name']:t for t in H}; by_f={t['name']:t for t in F}
loaded=[t for t in F if t['start_name']=='As loaded']
best=max(loaded,key=lambda t:t['score']); default=by_f['default_rep1']; half=by_f['half_steps_rep1']; wide=by_f['wide_first_rep1']
recommended_trials=[best]+[t for t in F if t['phase']=='start_sensitivity' and t['config']==best['config'] and t['params']['selection']==best['params']['selection']]
recommended_cap=400 if all(t['exhausted'] and t['accepted_moves']<400 for t in recommended_trials) else 800
schedule,base=opt.load_heavy_schedule(ROOT/opt.FILE_PATH,opt.SHEET_NAME)
horizon=fresh['horizon_index']
assert fresh['workbook_sha256']==hashlib.sha256((ROOT/opt.FILE_PATH).read_bytes()).hexdigest()
assert np.allclose(default['score'],by_h['baseline_default']['score'],rtol=0,atol=1e-12)
assert np.allclose(half['score'],by_h['rich_vns_with_halves']['score'],rtol=0,atol=1e-12)
cores=[]
for ac,events in schedule.items():
    due=events[0]['latest_start']+365*7
    assert next(e for e in events if e['check']=='D')['latest_start']==due
    cores.append((ac,due,due+126-70))
forced=np.array([sum(d<=day<end for _,d,end in cores) for day in range(horizon+1)])
peak=int(forced.max()); peakdays=np.flatnonzero(forced==peak)
selected=[c for c in cores if c[1]<=peakdays[0]<c[2]]
assert peak==6 and best['min']==18
improvement=(default['std']-best['std'])/default['std']*100
all_params={json.dumps(t['params'],sort_keys=True) for t in H+F}
all_scenarios={(json.dumps(t['params'],sort_keys=True),t.get('start_name','As loaded')) for t in H+F}
summary={'historical_runs':len(H),'fresh_runs':len(F),'distinct_parameter_sets':len(all_params),'distinct_setting_start_pairs':len(all_scenarios),'best_name':best['name'],'best_config':best['config'],'best_score':best['score'],'smoothness_improvement_percent':improvement,'forced_unavailable':peak,'peak_days':list(map(int,peakdays)),'all_historical_compliant':all(t['rule_compliant'] for t in H),'all_fresh_compliant':all(t['rule_compliant'] and t['independent_availability_verified'] for t in F)}
(WORK/'analysis-summary.json').write_text(json.dumps(summary,indent=2))

doc=Document(); sec=doc.sections[0]
sec.page_width=Inches(8.27);sec.page_height=Inches(11.69)
sec.left_margin=sec.right_margin=Inches(.78);sec.top_margin=Inches(.7);sec.bottom_margin=Inches(.65)
sec.footer_distance=Inches(.3)
for name in ['Normal','Title','Subtitle','Heading 1','Heading 2','Heading 3']:
    s=doc.styles[name];s.font.name='Calibri';s.font.color.rgb=RGBColor(0,0,0);s.paragraph_format.widow_control=True
for border in list(doc.styles.element.iter(qn('w:pBdr'))):border.getparent().remove(border)
doc.styles['Normal'].font.size=Pt(10.5);doc.styles['Normal'].paragraph_format.space_after=Pt(7);doc.styles['Normal'].paragraph_format.line_spacing=1.08
for name,size in [('Title',26),('Subtitle',12),('Heading 1',18),('Heading 2',12)]:
    s=doc.styles[name];s.font.size=Pt(size);s.paragraph_format.space_before=Pt(8);s.paragraph_format.space_after=Pt(8);s.paragraph_format.keep_with_next=True
doc.styles['Heading 1'].font.bold=True;doc.styles['Heading 2'].font.bold=True
footer=sec.footer.paragraphs[0];footer.alignment=WD_ALIGN_PARAGRAPH.RIGHT
r=footer.add_run('Priority based selection  |  ');r.font.size=Pt(8);r.font.color.rgb=RGBColor.from_string('666666')
field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE');footer._p.append(field)
doc.core_properties.title='Priority Based Selection Hyperparameter Analysis'
doc.core_properties.subject='Fleet maintenance neighbourhood selection convergence runtime and sensitivity'
doc.core_properties.author='David Long'
def p(text='',style=None):return doc.add_paragraph(text,style)
def h(text):doc.add_heading(text,2)
def page(title):doc.add_page_break();doc.add_heading(title,1)
def small(text):
    para=p(text)
    for r in para.runs:r.font.size=Pt(8.5);r.italic=True
    return para
def table(headers,rows,widths=None,font=9.2):
    t=doc.add_table(rows=1,cols=len(headers));t.alignment=WD_TABLE_ALIGNMENT.CENTER;t.autofit=False
    if widths is None:widths=[6.7/len(headers)]*len(headers)
    for c,w in zip(t.columns,widths):c.width=Inches(w)
    for c,txt in zip(t.rows[0].cells,headers):c.text=str(txt)
    for row in rows:
        for c,txt in zip(t.add_row().cells,row):c.text=str(txt)
    borders=OxmlElement('w:tblBorders')
    for edge in ['top','left','bottom','right','insideH','insideV']:
        e=OxmlElement('w:'+edge);e.set(qn('w:val'),'single');e.set(qn('w:sz'),'4');e.set(qn('w:color'),'D9D9D9');borders.append(e)
    t._tbl.tblPr.append(borders)
    for i,row in enumerate(t.rows):
        pr=row._tr.get_or_add_trPr();pr.append(OxmlElement('w:cantSplit'))
        if i==0:pr.append(OxmlElement('w:tblHeader'))
        for j,c in enumerate(row.cells):
            c.width=Inches(widths[j]);c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            cp=c._tc.get_or_add_tcPr();shade=OxmlElement('w:shd');shade.set(qn('w:fill'),'263746' if i==0 else ('F1F4F6' if i%2 else 'FFFFFF'));cp.append(shade)
            margins=OxmlElement('w:tcMar')
            for edge,val in [('top','60'),('bottom','60'),('left','75'),('right','75')]:
                e=OxmlElement('w:'+edge);e.set(qn('w:w'),val);e.set(qn('w:type'),'dxa');margins.append(e)
            cp.append(margins)
            for para in c.paragraphs:
                para.paragraph_format.space_after=Pt(0);para.paragraph_format.line_spacing=1.04
                para.alignment=WD_ALIGN_PARAGRAPH.LEFT if j==0 or len(para.text)>24 else WD_ALIGN_PARAGRAPH.CENTER
                for r in para.runs:r.font.size=Pt(font);r.bold=i==0;r.font.color.rgb=RGBColor.from_string('FFFFFF' if i==0 else '000000')
    p('').paragraph_format.space_after=Pt(1)
    return t
def f(x):return f'{x:.6f}'
def shortname(t):return t['name'].replace('_',' ')
COLORS=['#244C75','#AD5B27','#438078','#735A83']
from reportlab.lib.colors import HexColor
def chart(filename,series,xlim,ylim,xlabel,ylabel,xticks,yticks,scatter=False,height=260):
    width=640;left=65;bottom=48;pw=width-left-22;ph=height-bottom-35
    d=Drawing(width,height)
    def xp(x):return left+(x-xlim[0])/(xlim[1]-xlim[0])*pw
    def yp(y):return bottom+(y-ylim[0])/(ylim[1]-ylim[0])*ph
    for tick in yticks:
        yy=yp(tick);d.add(Line(left,yy,width-22,yy,strokeColor=HexColor('#DEDEDE'),strokeWidth=.6));d.add(String(left-9,yy-3,f'{tick:g}',fontSize=9,textAnchor='end'))
    for tick in xticks:d.add(String(xp(tick),bottom-16,f'{tick:g}',fontSize=9,textAnchor='middle'))
    d.add(Line(left,bottom,width-22,bottom,strokeColor=HexColor('#555555')))
    d.add(Line(left,bottom,left,height-35,strokeColor=HexColor('#555555')))
    d.add(String(left+pw/2,8,xlabel,fontSize=10,textAnchor='middle'))
    d.add(String(4,height-9,ylabel,fontSize=10))
    for i,(label,xs,ys) in enumerate(series):
        col=HexColor(COLORS[i%len(COLORS)]);coords=[]
        for x,y in zip(xs,ys):coords.extend([xp(float(x)),yp(float(y))])
        if not scatter:d.add(PolyLine(coords,strokeColor=col,strokeWidth=1.6))
        if scatter or len(xs)<20:
            for x,y in zip(xs,ys):d.add(Circle(xp(float(x)),yp(float(y)),2.6,fillColor=col,strokeColor=col))
        legendx=left+i*145;d.add(Line(legendx,height-15,legendx+14,height-15,strokeColor=col,strokeWidth=2));d.add(String(legendx+20,height-18,label,fontSize=9))
    path=WORK/(filename+'.pdf');renderPDF.drawToFile(d,str(path))
    pdf=pdfium.PdfDocument(path);out=WORK/(filename+'.png');pdf[0].render(scale=2.1).to_pil().save(out);pdf.close()
    doc.add_picture(str(out),width=Inches(6.65))
def shifts(t):return ' | '.join(','.join(f'{x:+d}' for x in s) for s in t['params']['shift_sets'])
def repeatstats(config):
    policy=best['params']['selection'] if config==best['config'] else 'first_level'
    items=[t for t in loaded if t['config']==config and t['params']['selection']==policy and t['phase']=='repeat']
    # A new leader's original refinement run is the first repeat in its timing set.
    if config==best['config'] and len(items)==2:items=[best]+items
    return items

p('Priority Based Selection Hyperparameter Analysis','Title')
p('Fleet maintenance optimisation','Subtitle');p('David Long\nUniversity of Queensland and Northrop Grumman\n13 October 2026')
h('Main finding')
p(f'The recommended settings for the as-loaded forecast are {best["config"].lower()} with {best["params"]["selection"].replace("_"," ")} selection and max_iters equal to {recommended_cap}. They achieve worst-day availability of {best["min"]} aircraft and smoothness of {best["std"]:.6f}, a {improvement:.2f} percent reduction in standard deviation from the default. Mean availability remains {best["avg"]:.6f} aircraft.')
p('Worst-day availability of 18 is globally maximal under the first D check rules for this forecast. Six aircraft must be in D maintenance together on six dates in May 2033. Tuning therefore improves the distribution of availability around an already optimal worst day. The best smoothness found is an empirical result; the study does not establish a global smoothness optimum.')
table(['Schedule','Worst day','Std deviation','Mean'],[['As loaded',15,f(fresh['original']['std']),f(fresh['original']['avg'])],['Default',18,f(default['std']),f(default['avg'])],['Previous tuned baseline',18,f(half['std']),f(half['avg'])],['Best in extended study',best['min'],f(best['std']),f(best['avg'])]],[2.8,1,1.4,1.5])
p(f'The analysis combines {len(H)} historical runs with {len(F)} fresh runs, covering {len(all_params)} distinct parameter sets and four start schedules. It examines iteration budgets, shift magnitudes, ladder structure, ladder order, selection policy, interactions, local refinements, runtime variability, and sensitivity to legal changes in the start schedule.')
small('Scope: vn_priority_based_selection.py on heavy_maintenance_forecast.xlsx, MxEvents. The simulated annealing and early D variants are outside this comparison.')

page('Algorithm and hyperparameters')
p('The baseline is a deterministic hill-climb using a sequence of neighbourhoods. For each neighbourhood, checks are ordered by the lowest availability during their active maintenance interval. The search scans priority levels from worst to best and stops at the first level containing an improving move. It then accepts the best move within that level and starts the next iteration at the first neighbourhood again.')
p('A proposal moves one check by an offset and rebuilds the aircraft\'s later checks. Rebuilds enforce the maintenance sequence and recalculate due dates and legal start windows. Positive offsets may move a check later within its legal window; no check may start after its current due date. A larger improving shift can cross over schedules that are worse, but the algorithm never accepts a worsening endpoint.')
table(['Setting','Default','Range analysed'],[['Iteration cap','400','5 to 800 plus full trace checkpoints'],['Shift magnitude in days','28 14 7 3','1 to 70'],['Neighbourhood structure','Four pairs','Pairs half offsets unions asymmetric sets'],['Neighbourhood order','Large to small','Reversed interleaved and refined ladders'],['Selection policy','First priority level','First level all levels first improving']],[1.75,1.45,3.5])
h('What is and is not tunable')
p('Only max_iters is exposed by vn_optimise. shift_sets is a local constant in that function, and first-priority-level selection is built into generate_neighbors. This study varies these algorithm settings in a separate driver. Temperature, cooling rate, tabu tenure, and a search random seed do not apply to this deterministic method.')
h('Objective and fixed inputs')
p('Scores are compared lexicographically as (minimum availability, negative standard deviation, mean availability). A gain of one aircraft on the worst day outranks any smoothness change. Standard deviation uses the population definition implemented by numpy.std. No weighted scalar objective is introduced.')
small('Maintenance rules are fixed domain inputs: A 14 day duration 365 day interval 42 day early allowance; B 28 730 70; C 84 1460 70; D 126 2920 70. Changing these rules would change the scheduling problem rather than tune the optimiser.')

page('Dataset and experimental design')
table(['Property','Value'],[['Aircraft and checks',f'{fresh["aircraft"]} aircraft and {fresh["checks"]} checks'],['Check counts','A 144 B 72 C 38 D 24'],['Observation dates',f'{fresh["base_date"]} to {fresh["last_date"]}'],['Forecast horizon','Day indices 0 to 4480 inclusive and 4481 observations'],['Total scheduled maintenance','10248 aircraft days'],['Starting score',f'15 and {fresh["original"]["std"]:.6f} and {fresh["original"]["avg"]:.6f}']],[2.7,4])
p('The historical evidence contains 74 run records and 72 distinct parameter sets. Duplicate baseline settings and different caps are retained for auditability. The historical file dates from 17 September 2026; the extension was executed on 13 October 2026. Fresh default and half-step runs reproduce the historical scores to numerical precision. The historical study did not store input hashes, so matching input statistics and reproduced endpoints establish consistency, rather than a verified historical byte identity.')
p('All comparisons use the same horizon fixed from the as-loaded schedule. The extension uses an 800-iteration safety cap to distinguish convergence from a budget limit. The final recommendation uses the measured convergence requirement. Historical caps, including 400, remain labelled with their own values in the full catalogue.')
p('Nine serial finalist repetitions measure elapsed time for the default, half steps, and wide-first ladder. Any new leading configuration receives further repetitions and start-schedule checks. Repeating a deterministic run measures runtime variability, not search quality variability. Three separately generated start schedules use fixed seeds 41, 42, and 43 solely to perturb legal starts.')
h('Validation')
p('The production default and study driver return the same complete schedule, score, and trace length. Every fresh endpoint passes complete maintenance-rule replay and a recomputation of daily availability. A separate boolean occupancy calculation verifies availability without the delta-array method. Every accepted score strictly improves, check counts and durations are conserved, and all 74 historical endpoints are marked rule-compliant. The nine existing baseline tests also pass in the analysis runtime.')
small(f'Runtime: Python {fresh["python"]} with NumPy {fresh["numpy"]} on {fresh["platform"]}. GUI plotting imports are omitted when loading the unchanged production scheduling functions into the bundled analysis runtime.')

page('Iteration budget and convergence')
series=[]
for t,label in [(default,'Default'),(half,'Half steps'),(wide,'Wide first')]:
    series.append((label,[x['iteration'] for x in t['trace']],[x['std'] for x in t['trace']]))
chart('convergence',series,(0,max(t['accepted_moves'] for t in [default,half,wide])),(1.35,2.0),'Accepted moves','Standard deviation',[0,80,160,240,320],[1.4,1.6,1.8,2.0])
small('Figure 1 Smoothness along the deterministic search path. Each point is an accepted move; the horizontal axis is not elapsed time.')
rows=[]
for cap in [5,10,20,40,80,160,200,280,319,400,800]:
    a=default['trace'][min(cap,default['accepted_moves'])];b=half['trace'][min(cap,half['accepted_moves'])]
    rows.append([cap,int(a['score'][0]),f(a['std']),int(b['score'][0]),f(b['std'])])
table(['Cap','Default min','Default std','Half min','Half std'],rows,[.7,1.3,1.6,1.3,1.8])
p('Checkpoint results are exact prefixes of the measured trajectories, rather than extra full runs. The default converges after 280 accepted moves and half steps after 319. An additional full pass detects exhaustion. Increasing a cap beyond natural convergence cannot improve an endpoint because the stopping rule still terminates the search.')
thresholds=[]
for t in [default,half,wide]:
    thresholds.append(f'{t["config"]} first reaches 18 after '+str(next(x['iteration'] for x in t['trace'] if x['min']==18))+' accepted moves')
small('; '.join(thresholds)+'. Caps that stop before that point sacrifice the primary objective. A 400 cap provides headroom for the as-loaded finalists; start variants may need more.')

page('Sensitivity to individual shift magnitudes')
singles=[t for t in H if t['phase']=='single_neighbourhood']
table(['Offset in days','Worst day','Std deviation','Moves','Exhausted'],[[abs(t['params']['shift_sets'][0][0]),int(t['min']),f(t['std']),t['accepted_moves'],'Yes' if t['exhausted'] else 'No'] for t in singles],[1.6,1.1,1.6,1,1.4])
p('A single plus-or-minus seven-day neighbourhood is the strongest converged single-magnitude setting in this sweep. It reaches 18 available aircraft but ends at standard deviation 1.470754, worse than the default and tuned ladders. A fourteen-day neighbourhood also reaches 18 quickly, after 188 accepted moves, but leaves standard deviation 1.474391.')
p('The one-, two-, and three-day searches hit the 400-move cap at worst-day availability 17. These results show inefficient progress within the tested budget; they do not establish that the small-step neighbourhoods would converge at 17. The five-day run also hits its cap, even though it reaches 18.')
p('Very wide single steps are poor on this workbook. The seventy-day setting accepts only two moves and converges at 15, while fifty-six days converges at 16. Wide jumps often lie outside legal windows or offer no improving endpoint. That explanation is consistent with the outcome but the historical study did not count rejected proposals, so it is not a measured rejection-rate result.')
h('Implication for tuning')
p('Retain multiple scales of movement. The strongest single neighbourhood does not replace a ladder, and larger offsets do not monotonically improve the result. Fine offsets are useful for later smoothness improvement when combined with earlier larger shifts.')
small('Source: historical single_neighbourhood phase. A No in Exhausted means the endpoint remains budget-limited and should not be described as a local optimum.')

page('Neighbourhood ladders and structure')
ladders=[t for t in H if t['phase']=='vns_ladder']
table(['Ladder','Worst day','Std deviation','Moves'],[[shortname(t).replace('vns ','').capitalize(),int(t['min']),f(t['std']),t['accepted_moves']] for t in sorted(ladders,key=lambda t:t['score'],reverse=True)],[3.1,1,1.6,1])
p('The dense ladder of plus-or-minus 28, 21, 14, 7, 3, and 1 days is the best ordinary paired ladder in the historical sweep. Its standard deviation is 1.396502. A geometric ladder of 32, 16, 8, 4, and 2 days improves from 1.404470 to 1.399695 when a one-day final neighbourhood is added.')
p('The coarse 70, 42, and 28 day ladder converges at worst-day availability 17. Two-level shortcuts reach 18 but lose smoothness. These observations support a mixture of coarse movement and fine finishing for this forecast. The exact ranking is path-dependent: the algorithm takes the first neighbourhood with an improvement, so adding or reordering a neighbourhood changes later schedules.')
h('Half offsets within each neighbourhood')
p('The previous tuned baseline uses four neighbourhoods: (-28, -14, 14, 28), (-14, -7, 7, 14), (-7, -3, 3, 7), and (-3, -1, 1, 3). The best improving move is selected among all offsets within the first improving priority level. This produces a different trajectory from trying each offset pair in a separate neighbourhood.')
p('A single union of the original offsets returns standard deviation 1.413656. Adding 21- and one-day offsets to that union reaches 1.406534. Both are worse than the half-step ladder. This is measured evidence for retaining the layered search structure on this input, rather than a general guarantee that layering always wins.')

page('Order sensitivity and neighbourhood ablations')
orders=[t for t in H if t['phase']=='order']+[t for t in F if t['phase']=='order']
table(['Order setting','Cap','Worst day','Std deviation','Exhausted'],[[shortname(t),t['params']['max_iters'],int(t['min']),f(t['std']),'Yes' if t['exhausted'] else 'No'] for t in orders],[2.7,.6,.8,1.4,1.2])
p('Historical small-first default and weekly ladders both stop at their 400 cap with worst-day availability 17. Their endpoints closely resemble the corresponding small single-step searches because the first small neighbourhood keeps providing improving moves. Fresh reverse-order experiments use a higher cap and therefore test whether the earlier result was merely insufficient budget.')
reverse_half=by_f['reverse_halves'];reverse_dense=by_f['reverse_dense']
p(f'Reversed half steps reaches {reverse_half["min"]} available aircraft and ends at standard deviation {reverse_half["std"]:.6f} after {reverse_half["accepted_moves"]} moves. The reversed dense ladder ends at {reverse_dense["min"]} with standard deviation {reverse_dense["std"]:.6f} after {reverse_dense["accepted_moves"]} moves. '+('It remains cap-limited, so its eventual endpoint is unresolved.' if not reverse_dense['exhausted'] else 'It reaches natural convergence under the tested neighbourhoods.'))
ab=[t for t in F if t['phase']=='ablation']
table(['Removed extra offsets','Worst day','Std deviation','Change from half steps'],[[t['config'],t['min'],f(t['std']),f(t['std']-half['std'])] for t in ab],[3,.8,1.4,1.5])
p('Each ablation restores one half-step neighbourhood to its original plus-or-minus pair, leaving the other three unchanged. A positive change in standard deviation means worse smoothness at the same worst day. This comparison isolates each local change in settings, although the resulting search path can differ from the first move onward.')
p('Ablations should be read together rather than added. The effect of removing offsets at one level depends on which other offsets remain available. The study measures these conditional effects and does not assign an independent contribution to each level.')
best_ablation=min(ab,key=lambda t:t['std'])
p(f'The strongest ablation is {best_ablation["config"].lower()}, with standard deviation {best_ablation["std"]:.6f}. Its change relative to the full half-step ladder is {best_ablation["std"]-half["std"]:+.6f}. This is a conditional comparison at the same worst-day availability, not evidence that the removed offsets are always beneficial or harmful.')
small('The original small-first claims apply to a 400 cap. Fresh reversed ladders are labelled separately with an 800 cap to avoid conflating ordering and stopping budget.')

page('Selection policy interactions')
rows=[]
for config in ['Default','Half steps','Dense pairs']:
    for policy in ['first_level','all_levels','first_improving']:
        candidates=[t for t in loaded if t['config']==config and t['params']['selection']==policy]
        t=candidates[0] if candidates else next(t for t in H if t['name']=='vns_dense')
        rows.append([config,policy.replace('_',' '),int(t['min']),f(t['std']),t['accepted_moves']])
table(['Neighbourhoods','Selection','Worst day','Std deviation','Moves'],rows,[1.4,1.7,.9,1.7,1])
p('First level returns all improving candidates in the first priority level that contains any, then takes the best. All levels takes the best improving candidate across every priority level for the active neighbourhood. First improving stops at the first improving candidate in priority level, start date, aircraft name, check index, and offset order. All three still require a strictly better lexicographic score.')
p('The interaction table tests each policy on three different neighbourhood structures. This avoids assuming that a selection rule measured on the default transfers unchanged to the tuned ladder. Ties in candidate scores retain the first candidate in the deterministic iteration order.')
bestpolicy=min([r for r in rows if r[2]==18],key=lambda r:float(r[3]))
p(f'The best endpoint in this interaction comparison uses {bestpolicy[0].lower()} with {bestpolicy[1]} selection and standard deviation {bestpolicy[3]}. The report recommends settings based on the complete extended comparison, including local refinements, rather than treating the selection comparison as a separate optimisation target.')
p('Priority sorting directs proposals towards checks that cover low-availability days, but it is a search heuristic. Even all-level selection preserves worst-day availability through the score rule. Any smoothness differences arise from different accepted paths, not from one policy allowing worse worst days.')
small('Fresh results use the same 800 safety cap. Historical default all-level selection ended at 1.415233; half-step all-level selection ended at 1.403198. Reproduced endpoints provide an additional consistency check.')

page('Local refinement and the recommended settings')
ref=[half]+[t for t in loaded if t['phase']=='local_refine']
table(['Setting','Worst day','Std deviation','Moves','Exhausted'],[[t['config'],t['min'],f(t['std']),t['accepted_moves'],'Yes' if t['exhausted'] else 'No'] for t in ref],[2.6,.8,1.4,.8,1.1])
p('The local refinement adds a two-day final pair, a forty-two-day first pair, or twenty-one-day offsets within the largest half-step neighbourhood. These are targeted nearby alternatives to the previous winner. Each changes one structural feature while preserving the maintenance rules and acceptance rule.')
h('Recommended configuration for the as loaded forecast')
p(f'{best["config"]} is the strongest lexicographic endpoint in the fresh as-loaded comparison. It accepts {best["accepted_moves"]} moves, '+('converges naturally' if best['exhausted'] else 'reaches the iteration cap')+f', and finishes at worst-day availability {best["min"]} with standard deviation {best["std"]:.6f}. Relative to the default, the absolute reduction in standard deviation is {default["std"]-best["std"]:.6f} aircraft, or {improvement:.2f} percent.')
for i,offsets in enumerate(best['params']['shift_sets'],1):p('Neighbourhood '+str(i)+'  '+', '.join(f'{x:+d}' for x in offsets)+' days')
p('Selection  '+best['params']['selection'].replace('_',' '))
p(f'Use max_iters equal to {recommended_cap}. The largest observed move count for this configuration across the as-loaded and three perturbed starts is {max(t["accepted_moves"] for t in recommended_trials)}. Stop naturally when a full pass finds no improving candidate and report whether the cap was reached. The fresh study used an 800 cap to check that the recommendation was not a budget-limited endpoint.')
p('The recommendation applies to this forecast and these rules. The controlled start checks on the next page assess sensitivity, while wider generalisation still requires forecasts from other fleet sizes and maintenance distributions.')

page('Sensitivity to the start schedule')
table(['Variant','Changed aircraft','Initial worst day','Initial std'],[['As loaded',0,15,f(fresh['original']['std'])]]+[[seed,v['changed_aircraft'],v['stats']['min'],f(v['stats']['std'])] for seed,v in fresh['start_variants'].items()],[1.9,1.7,1.5,1.6])
p('For each fixed seed, the analysis attempts one nonzero legal shift of up to fourteen days on a later check of every aircraft, then rebuilds that aircraft. The first due clocks, check sequence, check count, durations, fleet, and forecast horizon remain fixed. Later due dates respond to rebuilds under the same maintenance rules. These variants test local start sensitivity; they are not independent datasets.')
startrows=[]
for t in F:
    if t['phase']=='start_sensitivity':startrows.append([t['start_name'].replace('Perturbed ',''),t['config'],t['min'],f(t['std']),t['accepted_moves'],'Yes' if t['exhausted'] else 'No'])
table(['Seed','Configuration','Worst day','Std deviation','Moves','Exhausted'],startrows,[.5,2.1,.8,1.4,.8,1.1],font=8.8)
for config in sorted({t['config'] for t in F if t['phase']=='start_sensitivity'}):
    ts=[t for t in F if t['phase']=='start_sensitivity' and t['config']==config]
    p(f'{config} has median smoothness {statistics.median(t["std"] for t in ts):.6f}, ranging from {min(t["std"] for t in ts):.6f} to {max(t["std"] for t in ts):.6f} across the three perturbed starts. Worst-day availability ranges from {min(t["min"] for t in ts)} to {max(t["min"] for t in ts)}.')
winnercounts=Counter()
for seed in ['41','42','43']:
    candidates=[t for t in F if t['start_name']==f'Perturbed {seed}']
    winnercounts[max(candidates,key=lambda t:t['score'])['config']]+=1
p('The strongest measured configuration within each start variant is '+', '.join(f'{name} on {count} of three variants' for name,count in winnercounts.items())+'. A setting that wins on the as-loaded forecast need not win on every nearby start.')
small('Compare configurations within the same variant. A range across three nearby starts does not support a statistical confidence interval or a population robustness claim.')

page('Runtime and practical tradeoffs')
configs=list(dict.fromkeys(['Default','Half steps','Wide first',best['config']]))
rows=[]
for config in configs:
    ts=repeatstats(config)
    if len(ts)<3:continue
    times=[t['elapsed_seconds'] for t in ts]
    assert len({t['schedule_sha256'] for t in ts})==1
    rows.append([config,len(ts),f'{statistics.median(times):.2f}',f'{min(times):.2f} to {max(times):.2f}',f(ts[0]['std']),ts[0]['accepted_moves']])
table(['Configuration','Runs','Median seconds','Range seconds','Std deviation','Moves'],rows,[1.9,.5,1.1,1.3,1.2,.7])
p('Every repeated finalist returns an identical full schedule. Elapsed times measure the quiet study driver, including copying and trajectory statistics but excluding the final independent audits, file writes, plotting, and exports. The production default equivalence run is a validation step outside these timing samples. No experiment workers ran concurrently.')
p('Three timings give a useful local median and range, but host load and power management remain uncontrolled. Historical timing values span very different repetitions; for example, a geometric plus one-day refinement took 987 seconds while an earlier run took 95 seconds. The analysis therefore does not compare fresh medians with individual historical timings or rank all 74 historical settings by speed.')
baseline_time=statistics.median(t['elapsed_seconds'] for t in repeatstats('Default'))
for config in configs[1:]:
    ts=repeatstats(config)
    if len(ts)<3:continue
    median=statistics.median(t['elapsed_seconds'] for t in ts)
    p(f'{config} requires {median/baseline_time:.2f} times the default median runtime on this machine. Its standard deviation changes by {(default["std"]-ts[0]["std"])/default["std"]*100:.2f} percent at the same worst-day availability of 18.')
h('Choosing an operating point')
p('Use the best measured configuration for offline planning when its extra runtime is acceptable. Wide first is a modest change to the original ladder and offers a smaller smoothness gain. The default remains the frozen baseline for method comparisons. Select runtime and quality tradeoffs using reproduced timings rather than a historical outlier.')

page('Effects on daily availability and displacement')
series=[]
for t,label in [(fresh['original'],'As loaded'),(default,'Default'),(best,'Best tuned')]:
    xs=list(range(15,25));ys=[t['distribution'].get(str(x),0) for x in xs]
    series.append((label,xs,ys))
chart('availability_distribution',series,(15,24),(0,1600),'Available aircraft','Number of days',list(range(15,25)),[0,400,800,1200,1600])
small('Figure 2 Distribution of daily availability. Connected markers guide comparison between integer availability levels.')
table(['Schedule','Days below 19','Days at worst value','Std deviation'],[['As loaded',fresh['original']['days_below_19'],fresh['original']['worst_days'],f(fresh['original']['std'])],['Default',default['days_below_19'],default['worst_days'],f(default['std'])],['Best tuned',best['days_below_19'],best['worst_days'],f(best['std'])]],[2.2,1.5,1.6,1.4])
p('Worst-day frequency and standard deviation answer different questions. Minimising standard deviation at a fixed worst day can redistribute days among several availability levels. The score does not explicitly minimise the number of days below 19, so that count must be reviewed alongside smoothness rather than inferred from it.')
p(f'In this comparison, the default has {default["days_below_19"]} days below 19 and the best tuned schedule has {best["days_below_19"]}. The smoothness improvement therefore '+('comes from redistributing availability on other days while retaining the same count of low-availability days.' if default['days_below_19']==best['days_below_19'] else 'should be read alongside this separate change in low-availability-day frequency.'))
target=json.loads((WORK/(best['name']+'-schedule.json')).read_text());rows=[]
for check in 'ABCD':
    deltas=[new['start']-old['start'] for ac,events in schedule.items() for old,new in zip(events,target[ac]) if old['check']==check]
    rows.append([check,len(deltas),sum(d!=0 for d in deltas),min(deltas),max(deltas),f'{statistics.median(abs(d) for d in deltas):.1f}'])
table(['Check','Count','Moved','Earliest shift','Latest shift','Median absolute shift'],rows,[.6,.7,.8,1.4,1.4,1.8])
small('Start shifts are days relative to the as-loaded forecast and include rebuild effects. Positive values remain inside recalculated legal windows. There is no displacement limit in this baseline.')

page('Availability limits and operational scope')
p('The mean is fixed by conserved total downtime. Over 4481 days, the 278 checks require 10248 aircraft days of maintenance, so mean availability is 24 minus 10248 divided by 4481, or 21.713010489. Changing only legal starts cannot improve the mean when all checks stay inside the comparison horizon. The third score component therefore does not discriminate among these runs.')
h('Why worst day availability cannot exceed 18')
p('Each aircraft\'s first D due date is its first A due date plus 2555 days. A, B, and C checks do not reset the D due clock. The first D duration is 126 days and its early allowance is 70 days. A D check with due day d can start only between d minus 70 and d. Every legal placement occupies the forced interval from d to d plus 56, with the end excluded.')
table(['Aircraft','D due date','Last forced day'],[[ac,(base+__import__('pandas').Timedelta(days=d)).strftime('%d %b %Y'),(base+__import__('pandas').Timedelta(days=end-1)).strftime('%d %b %Y')] for ac,d,end in selected],[2.3,2.2,2.2])
p('These six forced intervals overlap on 20 to 25 May 2033. At least six aircraft must be unavailable on each of these dates, so at most 18 can be available. Since the tuned schedules achieve a worst day of 18 across the whole horizon, the primary objective is proven optimal for this forecast and the unchanged first D rules.')
p('This bound does not prove the best possible smoothness at minimum 18. A different neighbourhood or a search that accepts intermediate worsening schedules could still improve standard deviation. Reaching 19 would require changing a problem assumption, such as the D early allowance, first due clocks, or check duration, and must be treated as a separate scenario.')
h('Operational constraints')
p('The baseline validates maintenance rules, legal starts, non-overlap, and the horizon. It does not load operational_constraints.json and does not enforce busy periods, dependency groups, or displacement limits. A tuned schedule must be re-evaluated under any required operational constraints before it is used as an operational plan. Hyperparameter rankings under added constraints may differ.')

page('Implementation and reproducibility')
h('Applying the recommendation')
p('The production module and source workbook remain unchanged. To apply the recommended neighbourhoods in vn_optimise, replace the local shift_sets list with the neighbourhoods listed on the local refinement page. Preserve maintenance rebuilds and lexicographic acceptance. If the recommended policy is first level, no generate_neighbors change is needed; alternative policies require the corresponding study selection behaviour to be made explicit.')
p('Expose shift_sets as a function argument for repeatable comparison rather than editing a local constant for every experiment. Keep the frozen default as the default argument for baseline comparisons. Record the chosen cap, the accepted-move count, and whether the search stopped naturally. An 800 cap provides room for the observed start-schedule sensitivity experiments without requiring 800 moves.')
h('Reproducing the evidence')
p('Run .scratch/priority-hyperparameter-analysis/run_analysis.py with the Python runtime listed in the source record. The script resumes completed trials only when source workbook and production hashes match. To force a clean rerun, use a fresh study folder or preserve and rename the existing fresh-results.json before running again. The original study also skips existing names and must not be mistaken for a fresh execution when rerun against its stored file.')
p('fresh-results.json contains complete endpoint metrics, daily distributions, hashes, and accepted-move trajectories. Each fresh run also has a schedule JSON; start_41.json, start_42.json, and start_43.json preserve the sensitivity inputs. analysis-summary.json records combined counts and the selected configuration. The historical run table and fresh catalogue follow in the appendices.')
h('Limits of the recommendation')
p('This study covers one 24-aircraft forecast and three small legal perturbations of it. It does not test an independent fleet, a different horizon, larger data, changed maintenance rules, or operational constraints. The search is deterministic and the setting sweep is deliberately selected rather than exhaustive. The measured best is a tuned baseline for this problem, and future method comparisons should report comparable outcomes and budgets against it.')
p('The start-variant medians describe sensitivity across three fixed inputs. Timing medians describe three repeated executions of one deterministic input. Neither supports statistical significance claims. Report a cap-limited endpoint as incomplete convergence and avoid selecting settings solely from smoothness when their worst-day availability is lower.')

page('Historical run catalogue')
p('All 74 historical records are retained, including repeated settings and cap refinements. Standard deviation is in aircraft. Exhausted indicates that a full neighbourhood pass found no improving move. Seconds are historical observations and should not be compared directly with the fresh timing medians.')
table(['ID','Run name','Cap','Min','Std','Moves','End','Seconds'],[[f'H{i:02d}',shortname(t),t['params']['max_iters'],int(t['min']),f(t['std']),t['accepted_moves'],'Local' if t['exhausted'] else 'Cap',f'{t["elapsed_seconds"]:.1f}'] for i,t in enumerate(H,1)],[.42,2.38,.45,.35,.9,.5,.65,1.05],font=8.2)
small('Exact shift tuples and selection policy for each historical ID are stored under the matching run name in .scratch/priority-based-selection/hyperparameter-results.json. Local refers only to the tested neighbourhoods and policy, not a global optimum.')

page('Fresh run catalogue')
p('Fresh runs use an 800 cap. Endpoints pass rule replay and two availability recomputations. Start variants are identified by their generation seed. Repeated endpoint hashes confirm determinism. Seconds exclude validation and export work after the search.')
table(['ID','Run name','Start','Min','Std','Moves','End','Seconds'],[[f'F{i:02d}',shortname(t),'Loaded' if t['start_name']=='As loaded' else t['start_name'].replace('Perturbed ',''),t['min'],f(t['std']),t['accepted_moves'],'Local' if t['exhausted'] else 'Cap',f'{t["elapsed_seconds"]:.1f}'] for i,t in enumerate(F,1)],[.42,2.38,.5,.35,.9,.5,.65,1],font=8.2)

page('Source record and parameter definitions')
p('The report derives its numerical claims from local executable scheduling code and run records. Dates refer to the Brisbane study date. SHA256 hashes identify the current source bytes and stored result files.')
for label,path in [('Production algorithm','vn_priority_based_selection.py'),('Source workbook','heavy_maintenance_forecast.xlsx'),('Historical driver','.scratch/priority-based-selection/run_hyperparameter_study.py'),('Historical results','.scratch/priority-based-selection/hyperparameter-results.json'),('Fresh driver','.scratch/priority-hyperparameter-analysis/run_analysis.py'),('Fresh results','.scratch/priority-hyperparameter-analysis/fresh-results.json')]:
    h(label);p(path);digest=hashlib.sha256((ROOT/path).read_bytes()).hexdigest();small('SHA256 '+digest[:32]+'\n'+digest[32:])
h('Exact neighbourhood definitions')
table(['Configuration','Neighbourhoods in order'],[[name,' | '.join(', '.join(f'{x:+d}' for x in level) for level in offsets)] for name,offsets in analysis.CONFIGS.items()],[1.4,5.3],font=8.8)
small('Vertical bars separate neighbourhoods. Commas separate offsets evaluated within a neighbourhood. Each offset is measured in days. first_level is best improvement in the first improving priority level; all_levels is best improvement across levels; first_improving is the first improving candidate in deterministic order.')
page('Complete neighbourhood catalogue')
p('This catalogue makes the tested search structures explicit without requiring the raw JSON. Duplicate neighbourhood definitions are listed once. Caps and policies are identified in the run tables and selection comparison. Negative offsets move a check earlier; positive offsets move it later within its current legal window.')
seen=set();params=[]
for t in H+F:
    key=json.dumps(t['params']['shift_sets'])
    if key in seen:continue
    seen.add(key);params.append([shortname(t),shifts(t)])
table(['First run using structure','Offsets in days with levels separated by bars'],params,[2.25,4.45],font=8.5)
small('Structures that share offsets but differ in neighbourhood grouping or order are distinct. Within-neighbourhood offset order is also preserved because first-improving selection depends on it.')
out=ROOT/'docs/reports/2026-10-13-priority-based-selection-hyperparameter-analysis.docx'
out.parent.mkdir(parents=True,exist_ok=True);doc.save(out)
print(json.dumps(summary,indent=2));print('WROTE',out)
