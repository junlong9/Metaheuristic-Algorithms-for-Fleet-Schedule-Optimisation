"""Write the matched three-forecast report while retaining the earlier report."""
from pathlib import Path
import ast,sys,json,hashlib,math,statistics
from collections import Counter
import numpy as np
import pandas as pd
from docx import Document
from docx.shared import Inches,Pt,RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT,WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from reportlab.graphics.shapes import Drawing,String,Line,Rect,PolyLine,Circle
from reportlab.graphics import renderPDF
from reportlab.lib.colors import HexColor
import pypdfium2 as pdfium

ROOT=Path(__file__).resolve().parents[2];WORK=Path(__file__).resolve().parent
slugs=['original','moderate','wide'];labels={'original':'Original','moderate':'Moderate stagger','wide':'Wide stagger'}
files={s:WORK/f'three-dataset-{s}.json' for s in slugs}
parts={s:json.loads(files[s].read_text()) for s in slugs}
assert all(p['complete'] and p['stage']=='finalists' and p['datasets'][s]['done'] for s,p in parts.items())
configs=parts['original']['configs'];config_by={c['id']:c for c in configs}
assert all(p['configs']==configs and p['hashes']==parts['original']['hashes'] for p in parts.values())
trials=[t for p in parts.values() for t in p['trials']]
case={}
for s in slugs:
    for c in configs:
        options=[t for t in parts[s]['trials'] if t['config_id']==c['id'] and t['phase'] in ['sweep','budget_extension']]
        case[s,c['id']]=max(options,key=lambda t:t['score'])
best={s:max([case[s,c['id']] for c in configs],key=lambda t:t['score']) for s in slugs}
maxmin={s:best[s]['min'] for s in slugs}
eligible=[c for c in configs if all(case[s,c['id']]['min']==maxmin[s] and case[s,c['id']]['exhausted'] for s in slugs)]
def penalty(c,s):return (case[s,c['id']]['std']/best[s]['std']-1)*100
shared=min(eligible,key=lambda c:(max(penalty(c,s) for s in slugs),sum(penalty(c,s) for s in slugs))) if eligible else None
main=parts['original'];n=main['observations'];meta={s:parts[s]['datasets'][s] for s in slugs}
summary={'date':'2026-10-14','configurations':len(configs),'matched_sweep_cases':len(configs)*len(slugs),'total_run_records':len(trials),'best':{s:{k:best[s][k] for k in ['config_id','config_name','min','std','avg','accepted_moves','cap','exhausted']} for s in slugs},'shared_configuration':shared,'shared_worst_smoothness_penalty_percent':max(penalty(shared,s) for s in slugs) if shared else None}
(WORK/'three-dataset-summary.json').write_text(json.dumps(summary,indent=2))
doc=Document();sec=doc.sections[0];sec.page_width=Inches(8.27);sec.page_height=Inches(11.69)
sec.left_margin=sec.right_margin=Inches(.78);sec.top_margin=Inches(.7);sec.bottom_margin=Inches(.65);sec.footer_distance=Inches(.3)
for name in ['Normal','Title','Subtitle','Heading 1','Heading 2','Heading 3']:
    st=doc.styles[name];st.font.name='Calibri';st.font.color.rgb=RGBColor(0,0,0);st.paragraph_format.widow_control=True
for border in list(doc.styles.element.iter(qn('w:pBdr'))):border.getparent().remove(border)
doc.styles['Normal'].font.size=Pt(10.5);doc.styles['Normal'].paragraph_format.space_after=Pt(7);doc.styles['Normal'].paragraph_format.line_spacing=1.08
for name,size in [('Title',26),('Subtitle',12),('Heading 1',18),('Heading 2',12)]:
    st=doc.styles[name];st.font.size=Pt(size);st.paragraph_format.space_before=Pt(8);st.paragraph_format.space_after=Pt(8);st.paragraph_format.keep_with_next=True
doc.styles['Heading 1'].font.bold=doc.styles['Heading 2'].font.bold=True
footer=sec.footer.paragraphs[0];footer.alignment=WD_ALIGN_PARAGRAPH.RIGHT;r=footer.add_run('Priority based selection across three forecasts  |  ');r.font.size=Pt(8);r.font.color.rgb=RGBColor.from_string('666666')
field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE');footer._p.append(field)
doc.core_properties.title='Priority Based Selection Hyperparameter Analysis Across Three Forecasts';doc.core_properties.author='David Long'
doc.core_properties.subject='Original moderate stagger and wide stagger forecast comparison'
# Reuse the visually verified table and figure functions without running the earlier authoring script.
tree=ast.parse((WORK/'build_report.py').read_text());wanted={'p','h','page','small','table','f','chart','shifts'}
helpers=ast.Module(body=[node for node in tree.body if isinstance(node,ast.FunctionDef) and node.name in wanted],type_ignores=[])
COLORS=['#244C75','#AD5B27','#438078','#735A83'];exec(compile(helpers,'report_helpers','exec'),globals())
def v(t):return f'{t["min"]} / {t["std"]:.6f}'
def config_name(cid):
    name=config_by[cid]['name']
    return {'vns wide dense':'Wide dense ladder','vns dense':'Dense paired ladder','vns plus wide':'Default with a 42 day level'}.get(name,name.replace('refine ','').replace('vns ','').capitalize())
def samples(s,cid):return [t for t in parts[s]['trials'] if t['config_id']==cid and t['phase']=='timing_repeat']
def cap_for(ts):return max(400,max(int(math.ceil((t['accepted_moves']+1)/100)*100) for t in ts))

p('Priority Based Selection Hyperparameter Analysis Across Three Forecasts','Title')
p('Original moderate stagger and wide stagger datasets','Subtitle')
p('David Long\nUniversity of Queensland and Northrop Grumman\n14 October 2026')
h('Main finding')
if shared:
    p(f'The strongest broadly applicable setting for this maintenance-only benchmark is {shared["id"]}, {config_name(shared["id"])}. It reaches the highest measured worst-day availability on all three datasets. Its largest smoothness penalty relative to a separately tuned dataset winner is {summary["shared_worst_smoothness_penalty_percent"]:.2f} percent. Dataset-specific settings remain preferable when the final smoothness improvement matters.')
else:
    p('The study supports dataset-specific settings. No naturally converged tested configuration reaches the highest measured worst-day availability on all three inputs. A single universal recommendation would sacrifice the primary objective on at least one forecast.')
table(['Dataset','Initial minimum','Default minimum and std','Best minimum and std','Best setting'],[[labels[s],meta[s]['initial']['min'],v(case[s,'C01']),v(best[s]),best[s]['config_id']] for s in slugs],[1.6,.95,1.55,1.55,1.05])
p(f'The study crosses {len(configs)} neighbourhood and selection settings with the actual original, moderate stagger, and wide stagger workbooks, producing {len(configs)*3} matched sweep cases. It also records repeat timings, reference checks, and any finalist budget extensions. All {len(trials)} run records use the same scoring dates and preserve maintenance opening states.')
p('The moderate and wide datasets change initial maintenance ages and the required check mix. They are distinct problem instances. Their higher input availability is not an optimisation gain on the original forecast. All reported gains compare a configuration with the default on the same dataset.')
p(f'The previous half-step recommendation reaches {case["wide","C02"]["min"]} available aircraft on wide stagger, compared with {case["wide","C01"]["min"]} for the default. It therefore loses on the primary objective on that dataset and should not be adopted universally from the original-only result.')
small('This report extends the original-only study and replaces its recommendation as evidence for selecting settings across these three forecasts. It evaluates maintenance feasibility; operational constraints are audited separately.')

page('Datasets and maintenance opening states')
table(['Measure','Original','Moderate stagger','Wide stagger'],[['Aircraft']+[meta[s]['aircraft'] for s in slugs],['Supplied checks']+[meta[s]['checks'] for s in slugs],['A checks']+[meta[s]['check_counts'].get('A',0) for s in slugs],['B checks']+[meta[s]['check_counts'].get('B',0) for s in slugs],['C checks']+[meta[s]['check_counts'].get('C',0) for s in slugs],['D checks']+[meta[s]['check_counts'].get('D',0) for s in slugs],['Initial worst day']+[meta[s]['initial']['min'] for s in slugs],['Initial population std']+[f(meta[s]['initial']['std']) for s in slugs],['Initial mean']+[f(meta[s]['initial']['avg']) for s in slugs],['Checks already in progress']+[meta[s]['opening_checks_in_progress'] for s in slugs],['Full maintenance aircraft days']+[meta[s]['full_downtime'] for s in slugs]],[2.8,1.3,1.3,1.3])
p('The moderate input has three cohorts of eight aircraft with maintenance timelines advanced by one, three, or five model years. The wide input has six cohorts of four, advanced by one through six model years. A model year is 365 days. Cohort assignment uses the saved generation seed 42 and was not chosen by the optimiser.')
p('The aged fleets can start at A, B, or C within the repeating A B A C A B A D sequence. The existing opening-state adapter retains each aircraft\'s saved sequence index and A, B, C, and D due clocks. Every check that started before the benchmark start of 21 January 2026 is frozen. The original dataset retains its production rebuild assumption that its opening A cannot precede the first due date.')
p('Changing only FILE_PATH would incorrectly reconstruct the aged fleets as if all aircraft were at the original opening A position. The study instead reads .scratch/initial-maintenance-stagger/schedule-data.json and replays the correct opening state. Check counts, types, durations, and these opening clocks remain fixed within each dataset.')
small('Unchanged A B C D rules respectively use durations of 14 28 84 126 days, intervals of 365 730 1460 2920 days, and early allowances of 42 70 70 70 days. These are domain inputs, not search hyperparameters.')
small('All datasets have 24 aircraft, but their D counts and downtime differ. Mean availability should therefore not be compared across datasets as a pure measure of optimiser quality.')

page('Matched dates objective and validation')
p(f'All three inputs are scored from {main["forecast_start"]} through {main["forecast_end"]}, inclusive, for {n} daily observations. The internal coordinate origin is {main["coordinate_origin"]}, allowing checks that began before the scoring dates to retain their full intervals. The model horizon is day {main["model_horizon_index"]} and includes the full ends of supplied checks crossing the forecast boundary.')
p('The objective is the lexicographic tuple (minimum availability, negative population standard deviation, mean availability). A higher worst day outranks every smoothness gain. Scores are calculated with the same NumPy functions used by the production method, on the matched visible window. The algorithm never accepts a worse endpoint.')
p('The earlier report used 4481 observations through 28 April 2038. The matched study uses the published 4381-day forecast through 18 January 2038. Checks crossing either boundary remain in the model. Their starts can change the maintenance counted inside the visible window, so the visible mean can change even though full check durations are conserved. Standard deviations from the two horizons must not be compared directly.')
h('Verification of the study driver')
p('All settings preserve legal starts, no overlap, the maintenance sequence, and due resets. Each endpoint is checked by complete rule replay. Daily availability is independently recomputed by the delta-array method and by aircraft occupancy masks. Incremental availability must match both exactly. Past checks, opening states, check counts, and durations are preserved, and every accepted score strictly improves.')
p('Priority groups are cached within an iteration, unchanged maintenance intervals are omitted from availability updates, and clearly dominated proposals are screened conservatively. Retained proposals still receive the original exact score. Cached and reference neighbour lists match for every initial configuration and every final neighbourhood. Reference original-input runs also match complete trajectories and final schedule hashes after these performance changes.')
p('The live source was reorganised under vn_priority_based_selection.py during execution. The sweep processes had already loaded the preceding version. Serial repetitions therefore load byte-identical frozen algorithm and adapter snapshots whose hashes match all sweep records. This prevents a source-version change from being mixed into the timing or quality comparison.')
small('The 58-case sweeps run as independent processes. Their elapsed times are retained as audit observations but are excluded from performance comparisons. Finalist timing runs execute serially after all sweeps finish.')

page('Best configurations within each dataset')
for s in slugs:
    h(labels[s]);ranked=sorted([case[s,c['id']] for c in configs],key=lambda t:t['score'],reverse=True)[:5]
    table(['Setting','Minimum','Std deviation','Mean','Moves','End'],[[t['config_id'],t['min'],f(t['std']),f(t['avg']),t['accepted_moves'],'Local' if t['exhausted'] else 'Cap'] for t in ranked],[.7,.8,1.4,1.4,.9,1.5])
    baseline=case[s,'C01'];winner=best[s]
    p(f'{winner["config_id"]} uses {config_name(winner["config_id"])}. Relative to the default, worst-day availability changes from {baseline["min"]} to {winner["min"]}, and standard deviation changes from {baseline["std"]:.6f} to {winner["std"]:.6f}. '+(f'At the same worst day, this is a {(baseline["std"]-winner["std"])/baseline["std"]*100:.2f} percent smoothness improvement.' if baseline['min']==winner['min'] else 'The higher worst day is the decisive improvement under the lexicographic objective.'))
small('Local means natural exhaustion of the tested improving neighbourhoods. It does not mean globally optimal smoothness. Cap means the trajectory still has unresolved convergence.')

page('Individual shift magnitudes and budget limits')
single=[c for c in configs if len(c['params']['shift_sets'])==1 and len(c['params']['shift_sets'][0])==2 and c['params']['selection']=='first_level']
table(['Shift in days','Original minimum and std','Moderate minimum and std','Wide minimum and std'],[[abs(c['params']['shift_sets'][0][0])]+[v(case[s,c['id']])+(' cap' if not case[s,c['id']]['exhausted'] else '') for s in slugs] for c in single],[1.05,1.85,1.9,1.9])
p('This comparison holds selection fixed and uses one symmetric pair of offsets. The same shift does not necessarily rank equally on all inputs because the maintenance overlaps and feasible windows differ. A wide step can pass over worsening intermediate schedules, but its proposed endpoint still has to improve the full score.')
p('A cap-limited small-step run is not evidence that its neighbourhood converges at the reported worst day. The revised study uses an 800-move cap rather than assuming that a 400 cap resolves each ordering or step-size comparison. Exact prefix results on the convergence page show the consequences of smaller budgets without counting prefixes as independent runs.')
p('The variable-neighbourhood ladder remains a search setting, rather than a maintenance rule. The early allowances remain 42 days for A and 70 for B, C, and D. Temperature, cooling schedules, and a search random seed are outside this deterministic method.')

page('Ladder structure ordering and local changes')
selected=['C01','C02','C03','C04']+[c['id'] for c in configs if c['family'] in ['order','ablation','local_refine']]
selected=list(dict.fromkeys(selected))
table(['Setting','Structure','Original min and std','Moderate min and std','Wide min and std'],[[cid,config_name(cid)]+[v(case[s,cid]) for s in slugs] for cid in selected],[.5,2.2,1.3,1.35,1.35],font=8.4)
p('Half steps adds intermediate offsets inside each level. Dense pairs adds separate intermediate levels. These differ because the search restarts at its first level after every accepted move and takes the first neighbourhood containing an improvement. A union of offsets, a layered ladder, and a reversed ladder can follow different trajectories even when they use similar magnitudes.')
p('The ablations restore one half-step level to its original pair while retaining the other levels. Local refinements add two-day finishing, forty-two-day initial movement, or twenty-one-day offsets within the first level. Their effects are conditional on the remaining ladder and should not be added together as independent contributions.')
small('Complete ordered shift tuples for every setting are provided in the final catalogue. Some structures occur with more than one selection policy.')

page('Selection policy interactions across the forecasts')
lookup={json.dumps([c['params']['shift_sets'],c['params']['selection']]):c['id'] for c in configs}
rows=[]
for cid in ['C01','C02','C03']:
    for policy in ['first_level','all_levels','first_improving']:
        key=json.dumps([config_by[cid]['params']['shift_sets'],policy]);testid=lookup[key]
        rows.append([testid,config_by[cid]['name'],policy.replace('_',' ')]+[v(case[s,testid]) for s in slugs])
table(['ID','Ladder','Selection','Original','Moderate','Wide'],rows,[.4,1,1.4,1.3,1.3,1.3],font=8.6)
p('First level finds the first priority level containing any improving candidate and selects the best candidate in it. All levels selects the best improving candidate over every priority level in the active neighbourhood. First improving stops at the first improving candidate in deterministic order. Priority ties follow start date, aircraft name, check index, and offset order.')
p('All policies use the same acceptance rule, so all-level selection does not permit a lower worst day. Differences in final smoothness arise from different accepted paths. This crossed comparison tests whether a policy recommendation made on the original data also holds on the aged fleets.')
for s in slugs:
    candidates=[case[s,r[0]] for r in rows];winner=max(candidates,key=lambda t:t['score'])
    p(f'Within this interaction matrix, {labels[s].lower()} is strongest with {winner["config_id"]}, {config_by[winner["config_id"]]["params"]["selection"].replace("_"," ")}, reaching {winner["min"]} available aircraft and standard deviation {winner["std"]:.6f}.')

page('Convergence and iteration budget recommendations')
series=[]
for s in slugs:
    t=best[s];series.append((labels[s],[x['iteration'] for x in t['trace']],[-x['score'][1] for x in t['trace']]))
xmax=max(t['accepted_moves'] for t in best.values());ymin=min(t['std'] for t in best.values());ymax=max(meta[s]['initial']['std'] for s in slugs)
chart('three_dataset_convergence',series,(0,max(1,xmax)),(max(0,math.floor(ymin*10)/10-.1),math.ceil(ymax*10)/10),'Accepted moves','Population standard deviation',list(range(0,xmax+1,max(100,int(math.ceil(xmax/400))*100))),[round(ymin-.05,2),1.4,1.6,2.0],height=240)
small('Figure 1 Best measured settings on each dataset. The horizontal axis is accepted moves, not elapsed time. Lower smoothness across datasets reflects different starting check mixes as well as optimisation.')
rows=[]
for s in slugs:
    for cap in [20,40,80,160,200,400,800]:
        a=case[s,'C01'];b=best[s];aa=a['trace'][min(cap,a['accepted_moves'])]['score'];bb=b['trace'][min(cap,b['accepted_moves'])]['score']
        rows.append([labels[s].replace(' stagger',''),cap,f'{int(aa[0])} / {-aa[1]:.6f}',f'{int(bb[0])} / {-bb[1]:.6f}'])
table(['Dataset','Budget','Default minimum and std','Best setting minimum and std'],rows,[1.2,.7,2.4,2.4],font=8.6)
small('Checkpoint values are exact trajectory prefixes, not separate trials. A cap needs an extra iteration beyond the last accepted move to confirm natural exhaustion.')

page('Shared settings and dataset specific recommendations')
if shared:
    ordered=sorted(eligible,key=lambda c:(max(penalty(c,s) for s in slugs),sum(penalty(c,s) for s in slugs)))[:8]
    table(['Setting','Original std penalty','Moderate std penalty','Wide std penalty','Worst penalty'],[[c['id']]+[f'{penalty(c,s):.2f}%' for s in slugs]+[f'{max(penalty(c,s) for s in slugs):.2f}%'] for c in ordered],[.8,1.5,1.5,1.5,1.4])
    p(f'{shared["id"]} is the shared recommendation among naturally converged settings that retain each dataset\'s highest measured worst day. It minimises the largest relative smoothness penalty across the three inputs, using the average penalty to resolve a tie. It is a measured compromise on these forecasts, not a guarantee on a new fleet.')
else:p('No naturally converged configuration reaches all three dataset-specific worst-day maxima. Retain separate settings rather than trade away the primary objective for a universal parameter list.')
rows=[]
for s in slugs:
    t=best[s];rows.append([labels[s],t['config_id'],config_by[t['config_id']]['params']['selection'].replace('_',' '),t['accepted_moves'],cap_for([t])])
table(['Dataset','Setting','Selection','Moves to stop','Suggested cap'],rows,[1.7,.7,1.7,1.25,1.35])
if shared:
    p('Shared shift sets in order')
    for i,level in enumerate(shared['params']['shift_sets'],1):p(f'Level {i}  '+', '.join(f'{offset:+d}' for offset in level)+' days')
    p('Shared selection  '+shared['params']['selection'].replace('_',' '))
    p(f'Shared max_iters  {cap_for([case[s,shared["id"]] for s in slugs])}. This covers the measured accepted-move counts with room to detect exhaustion.')
p('The default remains the frozen baseline for method comparisons. A future annealing or basin-hopping comparison should face the tuned baseline on each input and use the same opening state, scoring dates, and operational assumptions.')

page('Serial runtime measurements')
rows=[]
for s in slugs:
    ids=list(dict.fromkeys(['C01','C02',best[s]['config_id']]+([shared['id']] if shared else [])))
    for cid in ids:
        ts=samples(s,cid)
        if len(ts)!=3:continue
        assert len({t['schedule_sha256'] for t in ts})==1
        assert all(t['schedule_sha256']==case[s,cid]['schedule_sha256'] for t in ts)
        secs=[t['elapsed_seconds'] for t in ts]
        rows.append([labels[s].replace(' stagger',''),cid,f'{statistics.median(secs):.2f}',f'{min(secs):.2f} to {max(secs):.2f}',case[s,cid]['accepted_moves']])
table(['Dataset','Setting','Median seconds','Range seconds','Accepted moves'],rows,[1.25,.7,1.5,1.85,1.4])
p('Each timing group contains three serial executions after the parallel sweep processes have finished. Repeated endpoints and full schedule hashes must agree. These repeats estimate elapsed-time variability for a deterministic search; they do not provide a distribution of schedule quality.')
p('Timing covers the search driver, copying, and trace recording. It excludes independent endpoint audits, neighbourhood equivalence checks after termination, operational audits, schedule exports, and report construction. The driver uses the same cached priority and exact-score screening implementation for every timing sample.')
p('Host load and power management remain uncontrolled, so the medians and ranges describe this local machine. Sweep timings and timings from the earlier original-only report are excluded from the comparative table. The archived sweeps include reference implementations and concurrent execution that would otherwise confound a speed comparison.')
small(f'Execution environment: Python {main["python"]}, NumPy {main["numpy"]}. No runtime claim is based on the fastest single run.')

page('Operational audits and visible boundary effects')
rows=[]
for s in slugs:
    for label,aud in [('Input',meta[s]['initial_operational_audit']),('Default',case[s,'C01']['operational_audit']),('Dataset best',best[s]['operational_audit'])]:
        rows.append([labels[s].replace(' stagger',''),label,aud['loose']['maximum_displacement']]+['Pass' if aud[mode]['passes'] else 'Fail' for mode in ['loose','medium','tight']])
table(['Dataset','Schedule','Max displacement days','Loose','Medium','Tight'],rows,[1.1,1.3,1.3,1,1,1])
p('The baseline study does not enforce operational_constraints.json during neighbour acceptance. The table audits unchanged loose, medium, and tight scenario requirements after optimisation. An endpoint can pass every maintenance rule and still fail a displacement limit, busy period, or dependency-group requirement. Such an endpoint is a benchmark result, not an approved operational plan.')
p('Displacement is measured against each dataset\'s own as-loaded checks and includes rebuild effects. Advancing hypothetical maintenance ages to create the moderate and wide inputs is not a legal move from the original dataset. The saved operational requirements are audited without relaxation.')
table(['Dataset','Input visible downtime','Best visible downtime','Fixed full downtime'],[[labels[s],round((24-meta[s]['initial']['avg'])*n),round((24-best[s]['avg'])*n),meta[s]['full_downtime']] for s in slugs],[1.8,1.6,1.6,1.7])
p('Visible maintenance days can change as legal starts move work across the forecast boundary. Full downtime remains fixed because durations and supplied check counts are preserved. A higher visible mean reflects changed exposure within the fixed window; it does not mean that required maintenance has been removed.')
small('Detailed busy-period and dependency breach counts and displacement checks are retained in each run record. Scenario dates and aircraft groups come from the saved unchanged operational configuration. Operationally constrained tuning should be a separately specified experiment with a chosen scenario enforced during search.')

page('Settings passing the operational audits')
rows=[]
for s in slugs:
    for mode in ['loose','medium','tight']:
        candidates=[dict(t,label=t['config_id'],max_displacement=t['operational_audit'][mode]['maximum_displacement']) for c in configs for t in [case[s,c['id']]] if t['operational_audit'][mode]['passes']]
        if meta[s]['initial_operational_audit'][mode]['passes']:
            initial=meta[s]['initial'];candidates.append(dict(initial,score=[initial['min'],-initial['std'],initial['avg']],label='Input',max_displacement=0))
        if candidates:
            winner=max(candidates,key=lambda t:t['score']);rows.append([labels[s].replace(' stagger',''),mode,winner['label'],winner['min'],f(winner['std']),winner['max_displacement']])
        else:rows.append([labels[s].replace(' stagger',''),mode,'None','NA','NA','NA'])
table(['Dataset','Scenario','Best audited setting','Worst day','Std deviation','Max displacement'],rows,[1.1,.9,1.5,.8,1.2,1.2],font=8.8)
p('This table filters the measured endpoints by the saved operational requirements and includes the unchanged input when it passes. It answers which tested snapshot is strongest among those that pass every scenario check. It is not a run of an optimiser that enforces those checks during search.')
p('Input means that the unchanged forecast outranks every passing optimisation endpoint in the tested set. None means no tested endpoint or unchanged input passes the scenario. It does not establish mathematical infeasibility of the scenario or show that a constrained search cannot find a better schedule.')
p('Post-search filtering can discard a strong trajectory after it has already taken an unacceptable move. Enforcing displacement limits, busy periods, and dependency groups during proposal generation would produce different accepted paths and potentially different best settings. The primary recommendations in this report apply to the maintenance-only benchmark until that separate constrained study is performed.')

page('Limits of the evidence and reproduction')
table(['Dataset','Forced first D overlap','Worst day upper bound','Best measured worst day'],[[labels[s],6 if s=='original' else 3,18 if s=='original' else 21,best[s]['min']] for s in slugs],[1.8,1.6,1.6,1.7])
p('Before the first forecast D check, smaller checks do not reset its opening D clock. With a 126-day duration and 70-day early allowance, every legal placement covers the 56-day interval beginning at that D due date. The peak overlap of these forced intervals bounds the worst day at 18 for the original and 21 for each stagger input. A measured result matching its bound proves optimality of the primary objective under those assumptions; otherwise the remaining gap is unresolved.')
p('These bounds do not establish globally optimal smoothness. The study is a broad discrete sweep on three related forecasts, including two synthetic maintenance-age variants, and uses each saved maintenance opening state. It does not establish performance on an independent real fleet, changed maintenance rules, other horizon lengths, or operationally constrained neighbourhoods.')
p('The study retains each supplied check list rather than generating additional check cycles when rebuilds change dates near the forecast boundary. This matches the fixed-list baseline. A planning implementation that regenerates the forecast and adds newly required checks would need a separate verification of the resulting schedules and parameter ranking.')
p('Run .scratch/priority-hyperparameter-analysis/run_three_datasets.py with --dataset original, moderate, or wide and --sweep-only for the independent sweep processes. After every sweep finishes, run finalise_three_datasets.py to obtain serial finalist timings and cap extensions, including the shared setting. The script resumes only against matching input, opening-state, frozen algorithm, frozen adapter, and configuration records.')
p('The per-dataset JSON files store scores, traces, audits, proposal counts, source hashes, and endpoint schedule hashes. Complete schedule JSON files are retained under three-dataset-schedules. The input workbooks, saved opening state, production optimiser, and operational configuration are not edited by the study.')
small('The earlier original-only Word report remains an archive for its longer scoring horizon and small legal start perturbations. Those perturbations are separate from the actual moderate and wide maintenance-age datasets analysed here.')

for s in slugs:
    page(labels[s]+' full result catalogue')
    p('Each setting uses an 800-move sweep cap. A listed 1600 cap indicates a subsequent finalist extension. Local means a full pass found no improving move. The table lists each structural setting once; repeated timings and reference runs remain in the raw records.')
    table(['ID','Min','Std deviation','Mean','Moves','Cap','End','Days below 19'],[[c['id'],case[s,c['id']]['min'],f(case[s,c['id']]['std']),f(case[s,c['id']]['avg']),case[s,c['id']]['accepted_moves'],case[s,c['id']]['cap'],'Local' if case[s,c['id']]['exhausted'] else 'Cap',case[s,c['id']]['days_below_19']] for c in configs],[.45,.5,1.2,1.2,.65,.55,.7,1.45],font=8.3)

page('Complete parameter catalogue')
p('Neighbourhoods are ordered from left to right and separated by a vertical bar. Offsets within a neighbourhood are in their tested order and measured in days. Selection is part of the structural setting, so identical offsets with different policies retain separate IDs.')
table(['ID','Selection','Ordered offsets in days'],[[c['id'],c['params']['selection'].replace('_',' '),' | '.join(','.join(f'{x:+d}' for x in level) for level in c['params']['shift_sets'])] for c in configs],[.4,1.1,5.2],font=8.4)

page('Source record')
rows=[]
for key,digest in main['hashes'].items():rows.append([{'production':'Frozen algorithm','adapter':'Frozen adapter'}.get(key,key),digest[:32]+'\n'+digest[32:]])
driver_digest=hashlib.sha256((WORK/'run_three_datasets.py').read_bytes()).hexdigest()
rows.append(['Final study driver',driver_digest[:32]+'\n'+driver_digest[32:]])
operations=json.loads((ROOT/'.scratch/initial-maintenance-stagger/schedule-data.json').read_text())
operations_digest=hashlib.sha256((ROOT/'operational_constraints.json').read_bytes()).hexdigest()
assert operations_digest==operations['constraints_sha256']
rows.append(['Operational rules',operations_digest[:32]+'\n'+operations_digest[32:]])
table(['Input or code','SHA256'],rows,[1.5,5.2],font=9)
p('Original input  heavy_maintenance_forecast.xlsx')
p('Moderate input  outputs/initial-stagger-20261013/heavy_maintenance_forecast_stagger_moderate.xlsx')
p('Wide input  outputs/initial-stagger-20261013/heavy_maintenance_forecast_stagger_wide.xlsx')
p('Opening state  .scratch/initial-maintenance-stagger/schedule-data.json')
p('Frozen opening adapter  .scratch/priority-hyperparameter-analysis/frozen-source/run_priority_optimise.py')
p('Frozen algorithm  .scratch/priority-hyperparameter-analysis/frozen-source/savn_priority_based_selection.py')
p('The frozen algorithm is the byte-identical CRLF form of Git revision f9459f4. The current live entry point is vn_priority_based_selection.py; the hashes above identify the version actually evaluated.')
p('Study driver  .scratch/priority-hyperparameter-analysis/run_three_datasets.py')
for s in slugs:
    digest=hashlib.sha256(files[s].read_bytes()).hexdigest();small(files[s].name+'\nSHA256 '+digest[:32]+'\n'+digest[32:])
out=ROOT/'docs/reports/2026-10-14-priority-hyperparameter-analysis-three-datasets.docx';doc.save(out)
print(json.dumps(summary,indent=2));print('WROTE',out)
