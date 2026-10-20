from pathlib import Path
import json
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT

ROOT=Path(__file__).resolve().parents[2]
ev=json.loads((ROOT/'.scratch/semester2-report/min19-evidence.json').read_text())
out=ROOT/'docs/reports/2026-10-12-min19-infeasibility-explained.docx'
doc=Document(); sec=doc.sections[0]
sec.page_width=Inches(8.27); sec.page_height=Inches(11.69)
sec.left_margin=sec.right_margin=Inches(.8)
sec.top_margin=Inches(.78); sec.bottom_margin=Inches(.7); sec.footer_distance=Inches(.3)
for border in list(doc.styles.element.iter(qn('w:pBdr'))): border.getparent().remove(border)
for n in ['Normal','Title','Subtitle','Heading 1','Heading 2']:
    st=doc.styles[n]; st.font.name='Calibri'; st.font.color.rgb=RGBColor(0,0,0)
    st.paragraph_format.widow_control=True
doc.styles['Normal'].font.size=Pt(11)
doc.styles['Normal'].paragraph_format.space_after=Pt(9)
doc.styles['Normal'].paragraph_format.line_spacing=1.1
for name,size in [('Title',25),('Subtitle',12),('Heading 1',18),('Heading 2',12)]:
    st=doc.styles[name]; st.font.size=Pt(size); st.paragraph_format.space_after=Pt(9)
    st.paragraph_format.space_before=Pt(10 if name!='Title' else 0)
    st.paragraph_format.keep_with_next=True
doc.styles['Heading 1'].font.bold=doc.styles['Heading 2'].font.bold=True
f=sec.footer.paragraphs[0]; f.alignment=WD_ALIGN_PARAGRAPH.RIGHT
r=f.add_run('Availability 19 feasibility report  |  '); r.font.size=Pt(8); r.font.color.rgb=RGBColor.from_string('666666')
field=OxmlElement('w:fldSimple'); field.set(qn('w:instr'),'PAGE'); f._p.append(field)
doc.core_properties.title='Why Minimum Fleet Availability 19 Is Infeasible'
doc.core_properties.author='David Long'
doc.core_properties.subject='Mandatory D check overlap in the current heavy maintenance forecast'

def p(text,style=None): return doc.add_paragraph(text,style)
def h(text): doc.add_heading(text,2)
def newpage(text): doc.add_page_break(); doc.add_heading(text,1)
def table(headers,rows,widths):
    t=doc.add_table(rows=1,cols=len(headers)); t.autofit=False
    for col,w in zip(t.columns,widths): col.width=Inches(w)
    for c,x in zip(t.rows[0].cells,headers): c.text=str(x)
    for row in rows:
        for c,x in zip(t.add_row().cells,row): c.text=str(x)
    borders=OxmlElement('w:tblBorders')
    for edge in ['top','left','bottom','right','insideH','insideV']:
        e=OxmlElement('w:'+edge); e.set(qn('w:val'),'single'); e.set(qn('w:color'),'D9D9D9'); e.set(qn('w:sz'),'4'); borders.append(e)
    t._tbl.tblPr.append(borders)
    for i,row in enumerate(t.rows):
        tr=row._tr.get_or_add_trPr(); tr.append(OxmlElement('w:cantSplit'))
        if i==0: tr.append(OxmlElement('w:tblHeader'))
        for j,c in enumerate(row.cells):
            c.width=Inches(widths[j]); c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            pr=c._tc.get_or_add_tcPr(); sh=OxmlElement('w:shd'); sh.set(qn('w:fill'),'263746' if i==0 else ('F2F4F6' if i%2 else 'FFFFFF')); pr.append(sh)
            mar=OxmlElement('w:tcMar')
            for side in ['top','bottom','left','right']:
                e=OxmlElement('w:'+side); e.set(qn('w:w'),'100'); e.set(qn('w:type'),'dxa'); mar.append(e)
            pr.append(mar)
            for para in c.paragraphs:
                para.paragraph_format.space_after=Pt(0); para.paragraph_format.line_spacing=1.05
                para.alignment=WD_ALIGN_PARAGRAPH.LEFT if j==0 else WD_ALIGN_PARAGRAPH.CENTER
                for r in para.runs:
                    r.font.size=Pt(10); r.font.bold=i==0; r.font.color.rgb=RGBColor.from_string('FFFFFF' if i==0 else '000000')
    p('').paragraph_format.space_after=Pt(0)
def reference(num,path,meaning):
    para=p('['+str(num)+'] ')
    rel=para.part.relate_to((ROOT/path).as_uri(),RT.HYPERLINK,is_external=True)
    link=OxmlElement('w:hyperlink'); link.set(qn('r:id'),rel)
    run=OxmlElement('w:r'); pr=OxmlElement('w:rPr'); col=OxmlElement('w:color'); col.set(qn('w:val'),'244C75'); pr.append(col); run.append(pr)
    txt=OxmlElement('w:t'); txt.text=path; run.append(txt); link.append(run); para._p.append(link)
    para.paragraph_format.space_after=Pt(3)
    x=p(meaning)
    for r in x.runs: r.font.size=Pt(9.5)

p('Why Minimum Fleet Availability 19 Is Infeasible','Title')
p('Evidence from the current heavy maintenance forecast','Subtitle')
p('David Long\nUniversity of Queensland and Northrop Grumman\n12 October 2026')
h('Finding')
p('The current forecast cannot maintain at least 19 available aircraft every day under the existing maintenance rules. Six aircraft must be in D checks simultaneously on 20 to 25 May 2033. With a fleet of 24 aircraft, that leaves at most 18 available. The priority-based method already achieves a worst-day availability of 18, so it reaches the best possible value for this objective on this instance. [1, 2, 3]')
h('What the target requires')
p('A minimum availability of 19 means that at least 19 of the 24 aircraft must be available on every day of the forecast. This permits no more than 5 aircraft in maintenance on any day. Just one day with 6 aircraft unavoidably in maintenance is enough to rule out the target. The current instance contains six such days.')
h('Why one D check has an unavoidable period')
p('A D check lasts 126 days. Its legal start window allows it to begin up to 70 days before its due date, but never after that date. Even if it starts at the earliest possible time, it still has 56 days remaining when its due date arrives, because 126 minus 70 equals 56.')
p('If the same check starts at the latest possible time, it starts on its due date. Every intermediate legal start also leaves the aircraft in maintenance from the due date through the following 56-day period. Changing the start can move downtime before or after that period, but cannot remove downtime inside it. I call this the check’s unavoidable interval.')
table(['Placement','Start','Finish boundary'],[['Earliest legal placement','70 days before due','56 days after due'],['Latest legal placement','On the due date','126 days after due'],['Overlap shared by every placement','On the due date','56 days after due']],[2.65,1.95,2.07])
p('The finish boundary is the first day the aircraft is available again. Thus the unavoidable interval includes its first day and excludes its finish boundary. It contains 56 maintenance days. [2]')

newpage('The six unavoidable intervals in the forecast')
p('The table gives the actual first-D due dates and the corresponding unavoidable maintenance days for the six aircraft. The dates come from the current local workbook. Each D check has a 126-day duration and a full 70-day start allowance. [1, 2]')
rows=[]
for c in ev['selected']:
    rows.append([c['aircraft'],c['earliest_date'],c['due_date'],c['due_date']+' to\n'+c['core_last_day_date']])
table(['Aircraft','Earliest D start','D due date','Unavoidable maintenance days'],rows,[1.2,1.35,1.3,2.82])
h('Where all six intervals overlap')
p('The common overlap begins on 20 May 2033, the latest of these six due dates. By that date, Aircraft 19 must have entered its D check, and the other five aircraft must already be in theirs.')
p('The overlap continues through 25 May 2033, the earliest last unavoidable day among the six aircraft. Aircraft 15 cannot finish its D check before 26 May, even if it starts on its earliest legal date of 20 January. Until that finish boundary, all six aircraft remain unavoidably in maintenance.')
table(['Date range','Aircraft that must be down','Maximum aircraft available'],[['20 to 25 May 2033','6','24 minus 6 equals 18']],[2.1,2.1,2.47])
h('A concrete example on 20 May')
p('For Aircraft 19, the earliest start is 11 March and the due date is 20 May. Starting on 11 March finishes on 15 July. Starting on 20 May finishes even later. Every legal start therefore keeps Aircraft 19 in maintenance on 20 May.')
p('For Aircraft 15, the earliest start is 20 January and the due date is 31 March. Starting on 20 January finishes on 26 May. Starting later also finishes later. Every legal start therefore keeps Aircraft 15 in maintenance on 20 May. The same reasoning applies to the other four aircraft in the table.')
p('On each of these six days, at least six distinct aircraft are unavailable. The target allows only five. No selection of legal start dates can satisfy both conditions.')

newpage('Why changing the search cannot reach 19')
h('The first D due dates remain fixed')
p('The optimiser sets each first D due date from that aircraft’s opening due date plus 2555 days, or 7 times 365 days. The opening due date is taken from the first check’s Latest_Start value in the workbook. It is fixed for this scheduling instance. [2]')
p('Moving earlier A, B, or C checks can change some later due windows, but those checks do not reset the D clock. In the replay rules, an A check resets A, a B check resets A and B, and a C check resets A, B, and C. Only a D check resets the D clock. Consequently the first D due dates in the table cannot move as a side effect of retiming earlier checks. [2]')
h('Every method faces the same bound')
p('Priority selection, simulated annealing, and basin-hopping all use these maintenance rules. They can explore different legal schedules, including temporary worse schedules, but every legal schedule still contains the six unavoidable D-check intervals. More trials, a higher temperature, a different seed, or a longer run cannot remove those mandatory days.')
p('The forecast therefore places an upper bound of 18 on worst-day availability. This bound follows from the constraints themselves. It does not depend on whether a particular search becomes trapped in a local optimum.')
h('The bound is attainable')
table(['Schedule','Worst day availability','Standard deviation','Mean availability'],[['As-loaded forecast','15','1.949767','21.713010'],['Saved priority-based schedule','18','1.413971','21.713010']],[2.2,1.5,1.5,1.47])
p('The saved priority-based schedule was reloaded and checked against the current maintenance replay rules. It is rule-compliant and its daily availability never falls below 18. Together, the upper bound and this feasible schedule establish that the optimal worst-day availability is exactly 18. This conclusion does not prove that its standard deviation is the best possible. Smoothness can still improve while worst-day availability stays at 18. [3]')
h('Average availability does not resolve the overlap')
p('The forecast contains 10248 aircraft days of maintenance across 4481 daily observations, giving an average of about 2.287 aircraft in maintenance and 21.713 aircraft available. A target of 19 would allow 5 aircraft down each day, so the total downtime is within that average allowance. The problem is the unavoidable concentration of six D checks on particular days. A good average cannot compensate for a day below the required minimum. [1, 2]')

newpage('Scope of the conclusion and supporting evidence')
h('Which assumptions would have to change')
p('Reaching 19 requires a change that removes this six-aircraft overlap or changes the fleet size. Possible changes to investigate include widening the permitted D start windows, shortening D-check duration, changing the opening due dates that determine the first D dates, or increasing the fleet. These are changes to the scheduling instance or maintenance assumptions, and would require separate justification.')
p('Removing this particular overlap would be necessary for a 24-aircraft fleet to reach 19, but would not by itself prove that 19 is achievable across the whole horizon. Other maintenance checks and constraints would still need to be considered in a new feasibility analysis.')
p('The argument applies to the current workbook and current rules. It does not claim that availability 19 is impossible for every fleet or every maintenance policy. It also does not rely on hangar-bay, crew, or parts capacity limits, which the baseline model does not enforce. Additional operational limits can restrict the feasible schedules further; they cannot remove this bound while the same first-D rules remain in force. [2]')
h('Verification')
p('The first-D windows for all 24 aircraft were checked against due dates reconstructed from their opening due dates. All matched. Counting unavoidable D occupancy across the forecast gave a peak of six aircraft on day indices 2676 through 2681, corresponding to 20 through 25 May 2033.')
p('For each of the six named aircraft, all 71 integer start dates in its 70-day D window were also checked against every day in the six-day overlap. All 2556 occupancy checks confirmed that the aircraft must be in maintenance. The saved optimised schedule provides the feasible result at 18. These checks reproduce the arithmetic; the unavoidable-interval argument explains why it holds for every legal schedule.')
h('Explanation for the presentation')
p('“A D check lasts 126 days but can only move 70 days early. That leaves 56 days when the aircraft must be in maintenance regardless of its start date. Six aircraft share that unavoidable period on 20 to 25 May 2033. With 24 aircraft, no legal schedule can have more than 18 available on those days. The priority method already achieves 18, so simulated annealing can improve smoothness but cannot raise the minimum to 19 under these rules.”')
h('Source files')
reference(1,'heavy_maintenance_forecast.xlsx','MxEvents contains the check dates, durations, and start windows used in the proof.')
reference(2,'vn_priority_based_selection.py','CHECK_RULES, _initial_due_dates, _update_due_dates, and rebuild_aircraft_schedule define the relevant constraints. compute_availability defines the daily count.')
reference(3,'optimised_heavy_maintenance_output.xlsx','Optimised_MxEvents contains the rule-compliant schedule achieving 18. Its availability and Summary sheets agree with the recalculated score.')

doc.save(out)
print(out)
