from pathlib import Path
import json
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT

ROOT = Path(__file__).resolve().parents[2]
folder = Path(__file__).parent
results = json.loads((folder / "comparison-results.json").read_text())
assert len(results["runs"]) == 4
priority = results["runs"][0]
candidate = results["runs"][1:]
out = ROOT / "docs/reports/2026-10-12-busy-period-20-tradeoff.docx"
doc = Document()
sec = doc.sections[0]
sec.page_width, sec.page_height = Inches(8.5), Inches(11)
sec.left_margin = sec.right_margin = Inches(.8)
sec.top_margin, sec.bottom_margin = Inches(.72), Inches(.65)
for border in list(doc.styles.element.iter(qn("w:pBdr"))):
    border.getparent().remove(border)
for name, size in [("Normal", 11), ("Title", 24), ("Subtitle", 12), ("Heading 1", 17), ("Heading 2", 12)]:
    style = doc.styles[name]
    style.font.name = "Calibri"
    style.font.size = Pt(size)
    style.font.color.rgb = RGBColor(0, 0, 0)
    style.paragraph_format.space_after = Pt(8)
    style.paragraph_format.line_spacing = 1.08
    style.paragraph_format.widow_control = True
    if name != "Normal":
        style.paragraph_format.keep_with_next = True
        style.paragraph_format.space_before = Pt(8)
doc.core_properties.title = "Availability 20 During a Busy Period"
doc.core_properties.subject = "Feasibility and availability trade-offs under a period-focused objective"
doc.core_properties.author = "David Long"
footer = sec.footer.paragraphs[0]
footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
footer.add_run("Semester 2 operational constraints addendum  |  ").font.size = Pt(8)
field = OxmlElement("w:fldSimple")
field.set(qn("w:instr"), "PAGE")
footer._p.append(field)

def p(text):
    return doc.add_paragraph(text)

def h(text):
    doc.add_heading(text, 2)

def newpage(title):
    doc.add_page_break()
    doc.add_heading(title, 1)

def table(headers, rows, widths):
    t = doc.add_table(rows=1, cols=len(headers))
    t.autofit = False
    for column, width in zip(t.columns, widths):
        column.width = Inches(width)
    for cell, text in zip(t.rows[0].cells, headers):
        cell.text = str(text)
    for row in rows:
        for cell, text in zip(t.add_row().cells, row):
            cell.text = str(text)
    borders = OxmlElement("w:tblBorders")
    for name in ("top", "left", "bottom", "right", "insideH", "insideV"):
        edge = OxmlElement("w:"+name)
        for k, v in [("val", "single"), ("color", "D9D9D9"), ("sz", "4")]:
            edge.set(qn("w:"+k), v)
        borders.append(edge)
    t._tbl.tblPr.append(borders)
    for i, row in enumerate(t.rows):
        row._tr.get_or_add_trPr().append(OxmlElement("w:cantSplit"))
        if i == 0:
            row._tr.get_or_add_trPr().append(OxmlElement("w:tblHeader"))
        for j, cell in enumerate(row.cells):
            cell.width = Inches(widths[j])
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            pr = cell._tc.get_or_add_tcPr()
            shade = OxmlElement("w:shd")
            shade.set(qn("w:fill"), "263746" if i == 0 else ("F2F4F6" if i % 2 else "FFFFFF"))
            pr.append(shade)
            margins = OxmlElement("w:tcMar")
            for side in ("top", "bottom", "left", "right"):
                edge = OxmlElement("w:"+side)
                edge.set(qn("w:w"), "95")
                edge.set(qn("w:type"), "dxa")
                margins.append(edge)
            pr.append(margins)
            for para in cell.paragraphs:
                para.paragraph_format.space_after = Pt(0)
                para.paragraph_format.line_spacing = 1.05
                para.alignment = WD_ALIGN_PARAGRAPH.LEFT if j == 0 else WD_ALIGN_PARAGRAPH.CENTER
                for run in para.runs:
                    run.font.size = Pt(10)
                    run.font.bold = i == 0
                    run.font.color.rgb = RGBColor.from_string("FFFFFF" if i == 0 else "000000")
    doc.add_paragraph("").paragraph_format.space_after = Pt(0)

def reference(path, explanation):
    para = doc.add_paragraph()
    rel = para.part.relate_to((ROOT/path).as_uri(), RT.HYPERLINK, is_external=True)
    link = OxmlElement("w:hyperlink")
    link.set(qn("r:id"), rel)
    run = OxmlElement("w:r")
    text = OxmlElement("w:t")
    text.text = path
    run.append(text)
    link.append(run)
    para._p.append(link)
    para.paragraph_format.space_after = Pt(2)
    para = p(explanation)
    for run in para.runs:
        run.font.size = Pt(9.5)

doc.add_paragraph("Availability 20 During a Busy Period", "Title")
doc.add_paragraph("Operational constraints and method comparison", "Subtitle")
p("David Long  |  Supervisor report and presentation preparation\n12 October 2026")
h("Finding")
p("I tested the two existing busy windows with a target of at least 20 available aircraft each day. The September to October window is achievable: an adapted priority-based search found a rule-compliant schedule with a busy-period minimum of 20 and a whole-forecast minimum of 18. The April to May window is infeasible under the current maintenance rules. The result depends on the dates being protected.")
table(["Busy period", "Minimum 20", "Evidence"], [
    ["27 April to 26 May 2033", "Infeasible", "Unavoidable D overlap caps availability at 18"],
    ["14 September to 13 October 2033", "Achieved", "Saved feasible priority schedule reaches 20 every day"]], [2.35, 1.15, 3.4])
p("Requiring 20 in both windows simultaneously is therefore infeasible. The successful September to October schedule protects that window on its own.")
h("Why the April to May window cannot reach 20")
p("A D check lasts 126 days and may start only 70 days early. Every legal placement therefore shares 56 unavoidable maintenance days from its due date. Six aircraft have overlapping unavoidable intervals on 20 to 25 May 2033, leaving at most 24 minus 6 equals 18 available. Five D checks are unavoidable on several surrounding days, so even 20 is ruled out there. Lower availability before or after the window cannot remove these mandatory days.")
h("What was allowed to change")
p("As requested, I allowed lower availability outside the September to October window. Maintenance durations, legal start windows, due resets, check sequence and non-overlap remained mandatory. I also retained the loose scenario’s 70-day displacement limit and its two dependency groups, each requiring at least two of four aircraft available every day. The original busy-period availability filters were replaced by the new target. No maintenance exemption was needed.")
p("These are the existing dummy operational dates and groups. The surrounding period was measured over the 30 days immediately before and after the target; availability changes across the rest of the forecast were recorded as well.")

newpage("The availability gained and the cost elsewhere")
p("The original score maximises the lowest availability anywhere in the forecast. It can reject a useful busy-period improvement if availability falls elsewhere. In this experiment I first maximised the busy-period minimum up to 20, then minimised shortage below 20, then improved availability outside the period and smoothness. The target was allowed to be unmet during repair, but a reported success had to meet it on every day.")
h("Measured priority result")
table(["Measure", "As loaded", "Earlier saved priority", "Busy period priority"], [
    ["Busy-period minimum", 15, 18, priority["period_minimum"]],
    ["Days below 20 in the period", 25, 30, priority["period_shortage_days"]],
    ["Whole-forecast minimum", 15, 18, priority["full_horizon_minimum"]],
    ["Minimum in 30 days before", 18, 19, priority["30_days_before_minimum"]],
    ["Minimum in 30 days after", 15, 19, priority["30_days_after_minimum"]],
    ["Forecast standard deviation", "1.949767", "1.413971", f'{priority["full_horizon_std"]:.6f}'],
    ["Mean availability", "21.713010", "21.713010", f'{priority["full_horizon_mean"]:.6f}']], [2.45, 1.15, 1.7, 1.6])
p("The adapted search first reached 20 after seven accepted priority moves, then continued improving the surrounding schedule. Its returned result used about 60 seconds and all checks passed maintenance replay, displacement and dependency validation. The largest displacement was 69 days.")
h("What sacrificing nearby availability means here")
p(f'The returned schedule gained {priority["net_period_aircraft_days_change_vs_saved_priority"]} aircraft-days inside the busy period compared with the earlier saved priority export, and lost exactly that amount outside it. In the 30 days before the busy period, worst-day availability fell from 19 to 18; the 30 days after it stayed at 19. The whole-forecast minimum still remained 18. Thus this example did not require reducing the global minimum to 17 or breaking maintenance rules.')
p(f'Compared with the as-loaded forecast, the busy period gained {priority["net_period_aircraft_days_change_vs_loaded"]} aircraft-days and the rest of the forecast lost the same net amount. Other dates also improved, so net loss should not be confused with the sum of all individual daily reductions. Total maintenance downtime and whole-forecast mean availability did not change.')
p("The earlier saved export is a historical maintenance-only reference, not the latest tuned result or a fair constrained competitor: it exceeds the current 70-day displacement limit on four checks. Its daily counts illustrate the trade-off; the new method comparison uses the same as-loaded start and retained constraints for both methods.")

newpage("Whether annealing adds value")
p("I also tested the existing basin-hopping candidate with the same adapted score, target-day priority, legal moves and retained constraints. Each run had a 60-second search budget, and the candidate used seeds 1, 2 and 3. Its warm-start and polish cap was reduced to 80 iterations to leave time for actual annealing escapes; the priority run retained its 400-iteration cap. This is an exploratory comparison under a time budget.")
rows = [["Priority", "20", str(priority["outside_minimum"]), f'{priority["full_horizon_std"]:.6f}', f'{priority["seconds"]:.1f}']]
for record in candidate:
    rows.append([f'Candidate seed {record["seed"]}', str(record["period_minimum"]), str(record["outside_minimum"]), f'{record["full_horizon_std"]:.6f}', f'{record["seconds"]:.1f}'])
table(["Method", "Period min", "Outside min", "Std dev", "Seconds"], rows, [2.05, 1.05, 1.2, 1.3, 1.3])
all_outside = [r["outside_minimum"] for r in candidate]
assert all(r["target_met"] and r["escape_moves"] > 0 for r in candidate)
median_outside = sorted(all_outside)[1]
median_std = sorted(r["full_horizon_std"] for r in candidate)[1]
p(f'All three candidate runs reached 20 and performed actual escape moves. Their median outside minimum was {median_outside}, and median full-forecast standard deviation was {median_std:.6f}. All reported plans passed maintenance and other retained operational checks. Every search reached its runtime budget, so these are best-so-far schedules rather than convergence claims.')
h("Recommendation for this case")
if max(all_outside) <= priority["outside_minimum"] and all(r["full_horizon_std"] >= priority["full_horizon_std"] for r in candidate):
    p("The adapted priority-based method is the practical choice for this tested window and budget. It already achieves 20 without reducing the whole-forecast minimum below 18; the candidate did not improve that result. The important change was giving the busy period priority in the objective. Annealing was not needed to find a feasible schedule here.")
else:
    p("Priority-based search already provides a feasible schedule at 20 while maintaining 18 elsewhere. The candidate results should be assessed against that complete outcome rather than solely against standard deviation. This exploratory test does not justify a general ranking across other busy periods or settings.")
p("For another feasible window, annealing may help if improving the period requires several coordinated moves through temporarily worse schedules. It cannot overcome an unavoidable maintenance overlap. A definitive method ranking would need more date windows and the usual 300-second budget, with settings and acceptable surrounding availability specified in advance.")
h("Relevant files")
reference(".scratch/busy-period-tradeoff/run_comparison.py", "Reproducible isolated experiment. Published baseline and candidate files were not edited.")
reference(".scratch/busy-period-tradeoff/comparison-results.json", "All reported measurements and experimental settings. Individual schedule JSON files are saved in the same folder.")
reference(".scratch/busy-period-tradeoff/window-check.json", "Daily unavoidable D overlap and window feasibility bounds.")
reference("operational_constraints.json", "Existing busy dates, displacement limits and dependency groups.")
doc.save(out)
print(out)
