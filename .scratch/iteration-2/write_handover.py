"""Write the Iteration 2 handover from study-results.json."""

import json
from pathlib import Path

from docx import Document
from docx.shared import Pt

ROOT = Path(__file__).resolve().parents[2]
RESULTS = Path(__file__).resolve().parent / "study-results.json"
MARKDOWN = ROOT / "docs" / "reports" / "2026-10-30-iteration-2-handover.md"
WORD = ROOT / "docs" / "reports" / "2026-10-30-iteration-2-handover.docx"


def fmt(row):
    if not row:
        return "not run"
    timeout = " timed out" if row.get("timed_out") else ""
    return (
        f"worst day {row['worst_day']}, "
        f"smoothness {row['smoothness']:.6f}, "
        f"{row['seconds']:.1f}s{timeout}"
    )


def table(headers, rows):
    lines = [
        "| " + " | ".join(headers) + " |",
        "| " + " | ".join("---" for _ in headers) + " |",
    ]
    for row in rows:
        lines.append("| " + " | ".join(str(cell) for cell in row) + " |")
    return "\n".join(lines)


def recommendation(results):
    decision = results.get("decision", {})
    judgement = decision.get("judgement")

    if not judgement:
        return (
            "The decision runs have not all finished, so this handover "
            "does not make the keep-or-drop call."
        )

    if judgement["keep"]:
        return (
            "Keep the candidate method. Its typical smoothness beats the "
            f"tuned baseline by at least 0.005 "
            f"(median smoothness {judgement['median_smoothness']:.6f}), "
            "every seed holds worst-day availability 18, and the typical "
            "run finishes within 5 minutes."
        )

    tuned = decision["tuned"]
    gain = tuned["smoothness"] - judgement["median_smoothness"]
    held = "yes" if judgement["holds_worst_day"] else "no"
    return (
        "Drop the candidate method. Its median smoothness was "
        f"{judgement['median_smoothness']:.6f}, a gain of {gain:.6f} "
        f"on the tuned baseline's {tuned['smoothness']:.6f}. The keep "
        "rule needs a gain of at least 0.005, worst-day availability 18 "
        "on every seed, and a typical run within 5 minutes. The "
        f"smoothness spread across seeds was {judgement['spread']:.6f}, "
        f"the typical runtime was {judgement['median_seconds']:.1f}s, "
        f"and every seed held worst-day availability 18: {held}. "
        "Recommend the tuned baseline instead: the frozen hill-climb, "
        "called with the half-step neighbourhoods (28 with 14, then 14 "
        "with 7, then 7 with 3, then 3 with 1, either direction) and a "
        "400-iteration cap. Leave the published defaults in the baseline "
        "module unchanged, and pass those neighbourhoods in from the caller."
    )


def build_markdown(results):
    decision = results.get("decision", {})
    sections = [
        "# Iteration 2 handover",
        "",
        "Semester 2 asked whether basin-hopping should replace the "
        "frozen priority hill-climb, and whether dummy operational "
        "limits can sit on top of it. This note records what was "
        "measured on the canonical heavy forecast. It does not fill in "
        "numbers for runs that did not finish.",
        "",
        "## Recommendation",
        "",
        recommendation(results),
        "",
        "Worst-day availability 19 is infeasible on this forecast. The "
        "17 September note shows six opening D checks force the floor "
        "to 18 or below, so later work should not spend itself hunting 19.",
        "",
        "Classic simulated annealing, from the earlier local chase, "
        "stayed at smoothness 1.413971 with worst-day availability 18. "
        "That is the published baseline smoothness. It is an ablation, "
        "not the method under test, and it was not re-run here.",
        "",
        "## How a worse move is accepted",
        "",
        "The execution-plan deck says a move that lowers worst-day "
        "availability is always rejected. The candidate method does not "
        "do that. Early on, a high temperature may accept that move, and "
        "it becomes rare as the search cools. The reported schedule is "
        "the best-so-far, so a dip is not what is handed back. Tightened "
        "acceptance is the fallback: it rejects a lower worst day before "
        "the temperature draw. The deck was left as presented.",
        "",
        "## Decision runs",
        "",
        "Same half-step neighbourhoods for the tuned baseline and the "
        "candidate method. The baseline is one run because it is "
        "deterministic. The candidate and the tightened fallback are the "
        "median of seeds 1, 2, and 3. A run that passes 5 minutes is stopped.",
        "",
    ]

    decision_rows = []
    if decision.get("tuned"):
        decision_rows.append(("Tuned baseline", fmt(decision["tuned"])))
    for row in decision.get("candidate", []):
        decision_rows.append((row["name"], fmt(row)))
    for row in decision.get("tightened", []):
        decision_rows.append((row["name"], fmt(row)))

    if decision_rows:
        sections.append(table(["Run", "Result"], decision_rows))
    else:
        sections.append("No decision runs have been recorded.")

    sections.extend(["", "## Start schedules", ""])
    sections.append(
        "The opening check on each aircraft is pinned to its due date "
        "by the maintenance replay, so a start schedule cannot move it. "
        "Bunched and spread move the next check, the first one the replay "
        "will place somewhere else in its start window, then rebuild. "
        "Due dates stay the same. These runs are unconstrained, so the "
        "start effect is not mixed with the operational limits. The "
        "candidate run from each start uses seed 1."
    )
    sections.append("")

    stagger_rows = []
    for row in results.get("stagger", []):
        stagger_rows.append((
            row["kind"],
            fmt(row["start"]),
            fmt(row["tuned"]),
            fmt(row["candidate"]),
        ))

    if stagger_rows:
        sections.append(table(
            ["Start", "Before search", "Tuned baseline", "Candidate, seed 1"],
            stagger_rows,
        ))
    else:
        sections.append("No start-schedule runs have been recorded.")

    sections.extend(["", "## Operational limits", ""])
    sections.append(
        "Limits live in operational_constraints.json. The workbook was "
        "not edited. Busy periods cover 27 April to 26 May 2033 and "
        "14 September to 13 October 2033, the busy stretches in the "
        "loaded forecast. Displacement limits are 70, 42, and 21 days "
        "from the as-loaded start. No group of four aircraft stays at "
        "three available on every day of the loaded plan, so loose and "
        "medium ask for two of four, which the loaded plan meets, and "
        "tight asks for three of four plus a busy-period floor of 16, "
        "which the loaded plan does not. Tight is reported infeasible. "
        "There is no repair."
    )
    sections.append("")

    constraint_rows = []
    for row in results.get("constraints", []):
        if row.get("feasible"):
            result = fmt(row.get("tuned"))
            breaches = len(row.get("tuned_breaches") or [])
            result = f"{result}; breaches after search: {breaches}"
        else:
            result = "infeasible at the loaded schedule; " + "; ".join(
                row.get("breaches") or []
            )
        constraint_rows.append((row["scenario"], result))

    if constraint_rows:
        sections.append(table(["Scenario", "Result"], constraint_rows))
    else:
        sections.append("No constraint runs have been recorded.")

    sections.extend(["", "## Settings sweep", ""])
    sections.append(
        "One setting at a time, on the tuned neighbourhoods, with no "
        "operational limits, seed 99. The hour cap stopped the list if "
        "it was still going. A second seed was used only for the best "
        "finished variant, and only if time remained. The decision-run "
        "defaults (temperature 1.5, 15 escape cycles, 30 plateau moves, "
        "400 polish iterations) were not repeated inside the hour. "
        "The extra seed of escape temperature 1.0 was stopped at 5 "
        "minutes during its opening polish, so that row is not a "
        "finished repeat of the setting."
    )
    sections.append("")

    sweep_rows = [
        (row["name"], fmt(row))
        for row in results.get("sweep", [])
    ]
    if sweep_rows:
        sections.append(table(["Variant", "Result"], sweep_rows))
    else:
        sections.append("No settings-sweep runs have been recorded.")

    sections.extend([
        "",
        "## Footnote on the cloud settings study",
        "",
        "An earlier 146-trial cloud study rebuilt the heavy forecast "
        "because the canonical workbook was not on the remote branch. "
        "That rebuilt file warm-started at worst-day availability 16, "
        "not 18. Those trials are not comparable with the numbers above.",
        "",
    ])
    return "\n".join(sections)


def write_word(markdown):
    document = Document()
    style = document.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(11)

    for block in markdown.split("\n\n"):
        if block.startswith("# "):
            document.add_heading(block[2:], level=0)
        elif block.startswith("## "):
            document.add_heading(block[3:], level=1)
        elif block.startswith("| "):
            rows = [
                [cell.strip() for cell in line.strip("|").split("|")]
                for line in block.splitlines()
                if not line.startswith("| ---")
            ]
            headers, body = rows[0], rows[1:]
            table_obj = document.add_table(rows=1 + len(body), cols=len(headers))
            table_obj.style = "Table Grid"
            for index, header in enumerate(headers):
                table_obj.rows[0].cells[index].text = header
            for row_index, row in enumerate(body, start=1):
                for col_index, value in enumerate(row):
                    table_obj.rows[row_index].cells[col_index].text = value
            document.add_paragraph()
        else:
            document.add_paragraph(block)

    document.save(WORD)


def main():
    results = json.loads(RESULTS.read_text(encoding="utf-8"))
    markdown = build_markdown(results)
    MARKDOWN.parent.mkdir(parents=True, exist_ok=True)
    MARKDOWN.write_text(markdown, encoding="utf-8")
    write_word(markdown)
    print(MARKDOWN)
    print(WORD)


if __name__ == "__main__":
    main()
